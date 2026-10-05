#!/usr/bin/env python
"""Resolve candidate PTM-use calls through their import origins."""

from __future__ import annotations

from mist.resources import DATA_DIR

import argparse
import ast
import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

import networkx as nx
from openpyxl import load_workbook
from mist.environment import project as analysis_project

from mist.io import count_values, display_output_path, iter_python_files, read_csv, write_csv
from mist.analysis.parse_cache import parse_python_files
from mist.rules.context_terms import CALL_DEFAULT_MODEL_ARGS, CALL_GENERIC_TERMINALS, build_call_term_config

GENERIC_TERMINALS = CALL_GENERIC_TERMINALS
DEFAULT_MODEL_ARGS = CALL_DEFAULT_MODEL_ARGS


@dataclass(frozen=True)
class LoaderRule:
    import_origin: str
    model_loader: str
    terminal_call: str
    chain_suffix: str
    model_args: frozenset[str]


@dataclass
class ImportBinding:
    alias: str
    import_origin: str = ""
    project_module: str = ""
    project_symbol: str = ""
    imported_name: str = ""
    line_number: int = 0


@dataclass
class AssignmentRecord:
    symbol: str
    target_name: str
    value: ast.AST
    line_number: int
    class_name: str = ""
    function_name: str = ""


@dataclass
class FunctionReturnRecord:
    symbol: str
    value: ast.AST
    line_number: int
    class_name: str = ""
    function_name: str = ""


@dataclass
class FunctionDefRecord:
    symbol: str
    class_name: str
    function_name: str
    params: list[str]
    default_values: dict[str, ast.AST]
    line_number: int
    is_property_like: bool = False


@dataclass(frozen=True)
class OriginResolution:
    origin: str = ""
    method: str = ""
    evidence: str = ""


@dataclass(frozen=True)
class JediOriginStats:
    enabled: bool = False
    origins_added: int = 0
    files_ok: int = 0
    files_failed: int = 0
    error: str = ""


@dataclass(frozen=True)
class LoaderStringResolution:
    origin: str
    loader_string: str
    evidence: str
    rules: tuple[LoaderRule, ...]


@dataclass(frozen=True)
class DynamicLoaderResolution:
    origin: str
    method: str
    evidence: str
    rules: tuple[LoaderRule, ...]


@dataclass
class CallRecord:
    file_path: str
    module: str
    class_name: str
    function_name: str
    node: ast.Call
    visible_chain: str
    chain_parts: list[str]
    keyword_args: list[str]
    has_model_payload: bool
    line_text: str


@dataclass
class FileInfo:
    path: Path
    rel_path: str
    module: str
    lines: list[str]
    imports: dict[str, ImportBinding]
    assignments: list[AssignmentRecord]
    returns: list[FunctionReturnRecord]
    functions: list[FunctionDefRecord]
    calls: list[CallRecord]
    model_payload_symbols: set[str]


def main(argv: Sequence[str] | None = None) -> int:
    """Build import-origin propagation evidence and write loader candidates."""
    args = parse_args(argv)
    repo = args.repo.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    configure_call_terms(args.reuse_codebook)
    rules = load_loader_rules(args.reuse_codebook)
    import_origins = {rule.import_origin for rule in rules if rule.import_origin}
    module_index = build_module_index(repo)
    occurrence_import_seed_keys = load_occurrence_import_origin_seed_keys(args.import_origins)
    files = parse_repo(
        repo,
        module_index,
        import_origins,
        syntax_cache=args.syntax_cache,
        occurrence_import_seed_keys=occurrence_import_seed_keys,
    )
    graph = nx.DiGraph()
    # First carry known origins through project symbols, then resolve dynamic string loaders.
    symbol_origins, origin_resolutions, jedi_origin_stats = propagate_origins(
        files,
        module_index,
        graph,
        repo=repo,
        import_origins=import_origins,
        enable_jedi=args.enable_jedi,
    )
    dynamic_loader_resolutions = propagate_dynamic_loaders(files, rules, graph)
    # Candidate matching happens after propagation so wrapper calls get a fair chance.
    candidates = find_loader_candidates(
        files=files,
        rules=rules,
        symbol_origins=symbol_origins,
        origin_resolutions=origin_resolutions,
        dynamic_loader_resolutions=dynamic_loader_resolutions,
        graph=graph,
        repo_full_name=args.repo_full_name,
        commit=args.commit,
    )

    write_csv(output_dir / "loader_candidates.csv", candidates)
    write_graph_edges(output_dir / "call_graph_edges.csv", graph)

    summary = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "repo": str(repo),
        "reuse_codebook": str(args.reuse_codebook),
        "import_origins": str(args.import_origins) if args.import_origins else "",
        "occurrence_import_origin_seed_keys": len(occurrence_import_seed_keys) if occurrence_import_seed_keys else 0,
        "syntax_cache": str(args.syntax_cache) if args.syntax_cache else "",
        "python_files_parsed": len(files),
        "symbols_with_origin": len(symbol_origins),
        "graph_nodes": graph.number_of_nodes(),
        "graph_edges": graph.number_of_edges(),
        "dynamic_loader_symbols": len(dynamic_loader_resolutions),
        "jedi_origin_edges": {
            "enabled": jedi_origin_stats.enabled,
            "origins_added": jedi_origin_stats.origins_added,
            "files_ok": jedi_origin_stats.files_ok,
            "files_failed": jedi_origin_stats.files_failed,
            "error": jedi_origin_stats.error,
        },
        "loader_candidates": len(candidates),
        "binding_eligible_counts": count_values(candidates, "binding_eligible"),
        "outputs": {
            "loader_candidates": display_output_path(output_dir / "loader_candidates.csv", resolve=True),
            "call_graph_edges": display_output_path(output_dir / "call_graph_edges.csv", resolve=True),
            "summary": display_output_path(output_dir / "call_summary.json", resolve=True),
        },
    }
    (output_dir / "call_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))
    return 0


def configure_call_terms(reuse_codebook: Path) -> None:
    """Load call analysis matching terms from config/codebook for this run."""

    global DEFAULT_MODEL_ARGS, GENERIC_TERMINALS
    term_config = build_call_term_config(reuse_codebook)
    DEFAULT_MODEL_ARGS = term_config.model_args
    GENERIC_TERMINALS = term_config.generic_terminals


def parse_args(argv: Sequence[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build import-origin anchored loader candidate graph.")
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument(
        "--reuse-codebook",
        type=Path,
        default=DATA_DIR / "reuse_rules.xlsx",
    )
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--repo-full-name", default="")
    parser.add_argument("--commit", default="")
    parser.add_argument("--syntax-cache", type=Path, help="Optional syntax cache directory.")
    parser.add_argument(
        "--import-origins",
        type=Path,
        help="Optional import_origin_occurrences.csv to use as explicit import-origin seeds.",
    )
    parser.add_argument(
        "--enable-jedi",
        action="store_true",
        help="Optionally add origin edges for symbols Jedi resolves to codebook import roots.",
    )
    return parser.parse_args(argv)


def load_occurrence_import_origin_seed_keys(path: Path | None) -> set[tuple[str, int, str]]:
    """Load exact import-origin seed locations emitted by identifier extraction."""
    if not path:
        return set()
    seeds: set[tuple[str, int, str]] = set()
    for row in read_csv(path, missing_ok=True):
        rel_path = row.get("file_path", "")
        line_number = int(row.get("line_number") or 0)
        origin = row.get("import_origin", "")
        decision = row.get("occurrence_decision", "include")
        if rel_path and line_number and origin and decision != "exclude":
            seeds.add((rel_path, line_number, origin))
    return seeds


def load_loader_rules(path: Path) -> list[LoaderRule]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["model_loading"]
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    pos = {header: headers.index(header) + 1 for header in headers if header}
    required = ["import origin", "model loader", "terminal_call", "chain_suffix", "model argument(s)"]
    missing = [header for header in required if header not in pos]
    if missing:
        raise ValueError(f"Missing codebook columns: {missing}")

    rules = []
    for row in range(2, ws.max_row + 1):
        model_loader = cell_text(ws.cell(row, pos["model loader"]).value)
        if not model_loader:
            continue
        model_args = parse_model_args(cell_text(ws.cell(row, pos["model argument(s)"]).value))
        rules.append(
            LoaderRule(
                import_origin=cell_text(ws.cell(row, pos["import origin"]).value),
                model_loader=model_loader,
                terminal_call=cell_text(ws.cell(row, pos["terminal_call"]).value),
                chain_suffix=cell_text(ws.cell(row, pos["chain_suffix"]).value),
                model_args=frozenset(model_args),
            )
        )
    return rules


def parse_model_args(value: str) -> set[str]:
    args = set(DEFAULT_MODEL_ARGS)
    for piece in value.replace("/", " ").replace(",", " ").replace(";", " ").split():
        piece = piece.strip()
        if not piece:
            continue
        args.add(piece.rsplit(".", 1)[-1].strip("[]"))
    return args


def build_module_index(repo: Path) -> dict[str, Path]:
    index = {}
    for path in iter_python_files(repo):
        rel = path.relative_to(repo)
        parts = list(rel.with_suffix("").parts)
        if parts[-1] == "__init__":
            parts = parts[:-1]
        module = ".".join(parts)
        if module:
            index[module] = path
        if parts and parts[0] in {"src", "lib"} and len(parts) > 1:
            index[".".join(parts[1:])] = path
    add_unique_suffix_module_aliases(index)
    return index


def add_unique_suffix_module_aliases(module_index: dict[str, Path]) -> None:
    """Allow script-style imports when a module suffix is unique in the repo."""
    suffix_refs: dict[str, set[Path]] = {}
    for module, path in list(module_index.items()):
        parts = module.split(".")
        for index in range(1, len(parts)):
            alias = ".".join(parts[index:])
            if not alias:
                continue
            suffix_refs.setdefault(alias, set()).add(path)
    for alias, paths in suffix_refs.items():
        if alias not in module_index and len(paths) == 1:
            module_index[alias] = next(iter(paths))


def parse_repo(
    repo: Path,
    module_index: dict[str, Path],
    import_origins: set[str],
    *,
    syntax_cache: Path | None = None,
    occurrence_import_seed_keys: set[tuple[str, int, str]] | None = None,
) -> dict[str, FileInfo]:
    files = {}
    path_to_modules = {}
    for module, path in module_index.items():
        path_to_modules.setdefault(path, []).append(module)

    for parsed in parse_python_files(repo, syntax_cache=syntax_cache):
        path = parsed.path
        rel_path = parsed.rel_path
        module = choose_module_name(path, path_to_modules)
        if parsed.tree is None:
            continue
        collector = FileCollector(
            rel_path=rel_path,
            module=module,
            lines=parsed.lines,
            module_index=module_index,
            import_origins=import_origins,
            occurrence_import_seed_keys=occurrence_import_seed_keys,
        )
        collector.visit(parsed.tree)
        files[rel_path] = FileInfo(
            path=path,
            rel_path=rel_path,
            module=module,
            lines=parsed.lines,
            imports=collector.imports,
            assignments=collector.assignments,
            returns=collector.returns,
            functions=collector.functions,
            calls=collector.calls,
            model_payload_symbols={
                assignment.symbol
                for assignment in collector.assignments
                if ast_contains_model_key(assignment.value)
            },
        )
    return files


class FileCollector(ast.NodeVisitor):
    def __init__(
        self,
        rel_path: str,
        module: str,
        lines: list[str],
        module_index: dict[str, Path],
        import_origins: set[str],
        occurrence_import_seed_keys: set[tuple[str, int, str]] | None = None,
    ) -> None:
        self.rel_path = rel_path
        self.module = module
        self.lines = lines
        self.module_index = module_index
        self.import_origins = import_origins
        self.occurrence_import_seed_keys = occurrence_import_seed_keys or set()
        self.imports: dict[str, ImportBinding] = {}
        self.assignments: list[AssignmentRecord] = []
        self.returns: list[FunctionReturnRecord] = []
        self.functions: list[FunctionDefRecord] = []
        self.calls: list[CallRecord] = []
        self.class_stack: list[str] = []
        self.function_stack: list[str] = []

    def visit_Import(self, node: ast.Import) -> None:
        for alias in node.names:
            root = alias.name.split(".", 1)[0]
            local = alias.asname or root
            if root in self.import_origins and self.occurrence_seed_allows(node.lineno, root):
                self.imports[local] = ImportBinding(
                    alias=local,
                    import_origin=root,
                    imported_name=alias.name,
                    line_number=node.lineno,
                )
            elif alias.name in self.module_index or root in self.module_index:
                self.imports[local] = ImportBinding(
                    alias=local,
                    project_module=alias.name if alias.name in self.module_index else root,
                    imported_name=alias.name,
                    line_number=node.lineno,
                )

    def visit_ImportFrom(self, node: ast.ImportFrom) -> None:
        if not node.module:
            return
        module = resolve_relative_module(self.module, node.module, node.level)
        root = module.split(".", 1)[0]
        for alias in node.names:
            local = alias.asname or alias.name
            if root in self.import_origins and self.occurrence_seed_allows(node.lineno, root):
                self.imports[local] = ImportBinding(
                    alias=local,
                    import_origin=root,
                    imported_name=f"{module}.{alias.name}",
                    line_number=node.lineno,
                )
            else:
                imported_module = f"{module}.{alias.name}"
                if imported_module in self.module_index:
                    self.imports[local] = ImportBinding(
                        alias=local,
                        project_module=imported_module,
                        imported_name=imported_module,
                        line_number=node.lineno,
                    )
                elif module in self.module_index:
                    self.imports[local] = ImportBinding(
                        alias=local,
                        project_module=module,
                        project_symbol=alias.name,
                        imported_name=f"{module}.{alias.name}",
                        line_number=node.lineno,
                    )

    def visit_FunctionDef(self, node: ast.FunctionDef) -> None:
        self.functions.append(
            FunctionDefRecord(
                symbol=qualified_symbol(self.module, self.class_stack, node.name),
                class_name=self.class_stack[-1] if self.class_stack else "",
                function_name=node.name,
                params=function_params(node),
                default_values=function_default_values(node),
                line_number=node.lineno,
                is_property_like=is_property_like_function(node),
            )
        )
        self.function_stack.append(node.name)
        self.generic_visit(node)
        self.function_stack.pop()

    def visit_AsyncFunctionDef(self, node: ast.AsyncFunctionDef) -> None:
        self.visit_FunctionDef(node)

    def visit_ClassDef(self, node: ast.ClassDef) -> None:
        self.class_stack.append(node.name)
        self.generic_visit(node)
        self.class_stack.pop()

    def visit_Assign(self, node: ast.Assign) -> None:
        for target in node.targets:
            symbol = self.symbol_for_target(target)
            if symbol:
                self.assignments.append(
                    AssignmentRecord(
                        symbol=symbol,
                        target_name=".".join(expr_chain(target)),
                        value=node.value,
                        line_number=node.lineno,
                        class_name=self.class_stack[-1] if self.class_stack else "",
                        function_name=self.function_stack[-1] if self.function_stack else "",
                    )
                )
        self.generic_visit(node)

    def visit_AnnAssign(self, node: ast.AnnAssign) -> None:
        if node.value is not None:
            symbol = self.symbol_for_target(node.target)
            if symbol:
                self.assignments.append(
                    AssignmentRecord(
                        symbol=symbol,
                        target_name=".".join(expr_chain(node.target)),
                        value=node.value,
                        line_number=node.lineno,
                        class_name=self.class_stack[-1] if self.class_stack else "",
                        function_name=self.function_stack[-1] if self.function_stack else "",
                    )
                )
        self.generic_visit(node)

    def visit_With(self, node: ast.With) -> None:
        """Treat context-manager bindings as assignments from their manager expressions."""
        for item in node.items:
            if item.optional_vars is None:
                continue
            symbol = self.symbol_for_target(item.optional_vars)
            if not symbol:
                continue
            self.assignments.append(
                AssignmentRecord(
                    symbol=symbol,
                    target_name=".".join(expr_chain(item.optional_vars)),
                    value=item.context_expr,
                    line_number=getattr(item.context_expr, "lineno", node.lineno),
                    class_name=self.class_stack[-1] if self.class_stack else "",
                    function_name=self.function_stack[-1] if self.function_stack else "",
                )
            )
        self.generic_visit(node)

    def visit_AsyncWith(self, node: ast.AsyncWith) -> None:
        self.visit_With(node)

    def visit_Return(self, node: ast.Return) -> None:
        if node.value is not None and self.function_stack:
            symbol = qualified_symbol(self.module, self.class_stack, self.function_stack[-1], "return")
            self.returns.append(
                FunctionReturnRecord(
                    symbol=symbol,
                    value=node.value,
                    line_number=node.lineno,
                    class_name=self.class_stack[-1] if self.class_stack else "",
                    function_name=self.function_stack[-1] if self.function_stack else "",
                )
            )
        self.generic_visit(node)

    def visit_Call(self, node: ast.Call) -> None:
        chain_parts = expr_chain(node.func)
        if chain_parts:
            self.calls.append(
                CallRecord(
                    file_path=self.rel_path,
                    module=self.module,
                    class_name=self.class_stack[-1] if self.class_stack else "",
                    function_name=self.function_stack[-1] if self.function_stack else "",
                    node=node,
                    visible_chain=".".join(chain_parts),
                    chain_parts=chain_parts,
                    keyword_args=sorted(keyword.arg for keyword in node.keywords if keyword.arg is not None),
                    has_model_payload=call_has_model_payload(node),
                    line_text=line_at(self.lines, node.lineno),
                )
            )
        self.generic_visit(node)

    def symbol_for_target(self, target: ast.AST) -> str:
        chain = expr_chain(target)
        if not chain:
            return ""
        if chain[0] == "self" and len(chain) >= 2 and self.class_stack:
            return f"{self.module}.{self.class_stack[-1]}.{'.'.join(chain[1:])}"
        if len(chain) == 1:
            return qualified_symbol(self.module, self.class_stack, chain[0])
        return qualified_symbol(self.module, self.class_stack, ".".join(chain))

    def occurrence_seed_allows(self, line_number: int, root: str) -> bool:
        if not self.occurrence_import_seed_keys:
            return True
        return (self.rel_path, line_number, root) in self.occurrence_import_seed_keys


def propagate_origins(
    files: dict[str, FileInfo],
    module_index: dict[str, Path],
    graph: nx.DiGraph,
    *,
    repo: Path,
    import_origins: set[str],
    enable_jedi: bool = False,
) -> tuple[dict[str, str], dict[str, OriginResolution], JediOriginStats]:
    origins: dict[str, str] = {}
    origin_resolutions: dict[str, OriginResolution] = {}

    for file in files.values():
        file_node = f"file:{file.rel_path}"
        graph.add_node(file_node, kind="file", label=file.rel_path)
        for binding in file.imports.values():
            alias_symbol = qualified_symbol(file.module, [], binding.alias)
            if binding.import_origin:
                origins[alias_symbol] = binding.import_origin
                origin_resolutions[alias_symbol] = OriginResolution(
                    origin=binding.import_origin,
                    method="import_origin_alias",
                    evidence=binding.imported_name or binding.alias,
                )
                origin_node = f"origin:{binding.import_origin}"
                alias_node = f"symbol:{alias_symbol}"
                graph.add_node(origin_node, kind="import_origin", label=binding.import_origin)
                graph.add_node(alias_node, kind="symbol", label=alias_symbol)
                graph.add_edge(file_node, origin_node, kind="imports_origin", line=binding.line_number)
                graph.add_edge(origin_node, alias_node, kind="binds_alias", line=binding.line_number)

    jedi_origin_stats = add_jedi_origin_edges(
        repo,
        files,
        import_origins,
        origins,
        origin_resolutions,
        graph,
        enabled=enable_jedi,
    )

    changed = True
    for _ in range(8):
        if not changed:
            break
        changed = False
        for file in files.values():
            env = file_env(file, origins)
            for function in file.functions:
                for param, default in function.default_values.items():
                    resolution = expr_origin_resolution(
                        default,
                        file,
                        env,
                        origins,
                        origin_resolutions,
                        function.class_name,
                        function.function_name,
                    )
                    if not resolution.origin:
                        continue
                    for symbol in parameter_symbols(
                        file.module,
                        function.class_name,
                        function.function_name,
                        param,
                    ):
                        if origins.get(symbol) != resolution.origin:
                            origins[symbol] = resolution.origin
                            origin_resolutions[symbol] = OriginResolution(
                                resolution.origin,
                                f"function_default_origin:{resolution.method}",
                                resolution.evidence,
                            )
                            changed = True
                            add_origin_edge(
                                graph,
                                resolution.origin,
                                symbol,
                                f"function_default_origin:{resolution.method}",
                                function.line_number,
                            )
            for assignment in file.assignments:
                resolution = expr_origin_resolution(
                    assignment.value,
                    file,
                    env,
                    origins,
                    origin_resolutions,
                    assignment.class_name,
                    assignment.function_name,
                )
                if resolution.origin and origins.get(assignment.symbol) != resolution.origin:
                    origins[assignment.symbol] = resolution.origin
                    origin_resolutions[assignment.symbol] = resolution
                    changed = True
                    add_origin_edge(
                        graph,
                        resolution.origin,
                        assignment.symbol,
                        f"assignment_from_origin:{resolution.method}",
                        assignment.line_number,
                    )
            for ret in file.returns:
                resolution = expr_origin_resolution(
                    ret.value,
                    file,
                    env,
                    origins,
                    origin_resolutions,
                    ret.class_name,
                    ret.function_name,
                )
                if resolution.origin and origins.get(ret.symbol) != resolution.origin:
                    origins[ret.symbol] = resolution.origin
                    origin_resolutions[ret.symbol] = resolution
                    changed = True
                    add_origin_edge(
                        graph,
                        resolution.origin,
                        ret.symbol,
                        f"function_returns_origin:{resolution.method}",
                        ret.line_number,
                    )
            for function in file.functions:
                if not function.is_property_like or not function.class_name:
                    continue
                return_symbol = qualified_symbol(
                    file.module,
                    [function.class_name],
                    function.function_name,
                    "return",
                )
                if return_symbol not in origins:
                    continue
                property_symbol = qualified_symbol(file.module, [function.class_name], function.function_name)
                if origins.get(property_symbol) == origins[return_symbol]:
                    continue
                resolution = origin_resolutions.get(
                    return_symbol,
                    OriginResolution(origins[return_symbol], "property_return_origin", function.function_name),
                )
                origins[property_symbol] = resolution.origin
                origin_resolutions[property_symbol] = OriginResolution(
                    resolution.origin,
                    f"property_value_origin:{resolution.method}",
                    resolution.evidence,
                )
                changed = True
                add_origin_edge(
                    graph,
                    resolution.origin,
                    property_symbol,
                    f"property_value_origin:{resolution.method}",
                    function.line_number,
                )
    return origins, origin_resolutions, jedi_origin_stats


def add_jedi_origin_edges(
    repo: Path,
    files: dict[str, FileInfo],
    import_origins: set[str],
    origins: dict[str, str],
    origin_resolutions: dict[str, OriginResolution],
    graph: nx.DiGraph,
    *,
    enabled: bool,
) -> JediOriginStats:
    """Let Jedi add origin edges only when it resolves to a codebook import root."""
    if not enabled:
        return JediOriginStats(enabled=False)
    try:
        import jedi  # type: ignore
    except Exception as exc:
        return JediOriginStats(enabled=True, error=f"jedi_unavailable:{type(exc).__name__}:{exc}")

    try:
        project = analysis_project(repo, jedi_sys_paths(repo))
    except Exception as exc:
        return JediOriginStats(enabled=True, error=f"jedi_project_error:{type(exc).__name__}:{exc}")

    origins_added = 0
    files_ok = 0
    files_failed = 0
    last_error = ""
    for file in files.values():
        try:
            script = jedi.Script(path=str(file.path), project=project)
            for symbol, node in jedi_origin_probe_nodes(file):
                if symbol in origins:
                    continue
                resolved_origin, evidence = jedi_resolved_import_origin(script, node, import_origins)
                if not resolved_origin:
                    continue
                origins[symbol] = resolved_origin
                origin_resolutions[symbol] = OriginResolution(
                    resolved_origin,
                    "jedi_resolved_import_origin",
                    evidence,
                )
                add_origin_edge(
                    graph,
                    resolved_origin,
                    symbol,
                    "jedi_resolved_import_origin",
                    getattr(node, "lineno", 0),
                )
                origins_added += 1
            files_ok += 1
        except Exception as exc:
            files_failed += 1
            last_error = f"{type(exc).__name__}:{exc}"
    return JediOriginStats(
        enabled=True,
        origins_added=origins_added,
        files_ok=files_ok,
        files_failed=files_failed,
        error=last_error,
    )


def jedi_origin_probe_nodes(file: FileInfo) -> list[tuple[str, ast.AST]]:
    probes: list[tuple[str, ast.AST]] = []
    for assignment in file.assignments:
        for node in origin_probe_exprs(assignment.value):
            if assignment.symbol:
                probes.append((assignment.symbol, node))
            chain = expr_chain(node)
            if chain:
                probes.append((qualified_symbol(file.module, [], chain[0]), node))
    for function in file.functions:
        for param, default in function.default_values.items():
            for node in origin_probe_exprs(default):
                for symbol in parameter_symbols(
                    file.module,
                    function.class_name,
                    function.function_name,
                    param,
                ):
                    probes.append((symbol, node))
    return unique_probe_nodes(probes)


def origin_probe_exprs(node: ast.AST) -> list[ast.AST]:
    if isinstance(node, ast.Call):
        return [node.func]
    if isinstance(node, (ast.Name, ast.Attribute)):
        return [node]
    return []


def unique_probe_nodes(probes: Iterable[tuple[str, ast.AST]]) -> list[tuple[str, ast.AST]]:
    seen = set()
    output = []
    for symbol, node in probes:
        key = (symbol, getattr(node, "lineno", 0), getattr(node, "col_offset", 0))
        if key in seen:
            continue
        seen.add(key)
        output.append((symbol, node))
    return output


def jedi_resolved_import_origin(
    script: Any,
    node: ast.AST,
    import_origins: set[str],
) -> tuple[str, str]:
    line = getattr(node, "lineno", 0)
    if not line:
        return "", ""
    column = jedi_lookup_column(node)
    try:
        definitions = script.goto(line=line, column=column, follow_imports=True, follow_builtin_imports=False)
    except Exception:
        return "", ""
    for definition in definitions:
        full_name = getattr(definition, "full_name", None) or getattr(definition, "module_name", "")
        origin = import_origin_from_full_name(str(full_name or ""), import_origins)
        if origin:
            return origin, str(full_name)
    return "", ""


def import_origin_from_full_name(full_name: str, import_origins: set[str]) -> str:
    if not full_name:
        return ""
    for origin in sorted(import_origins, key=len, reverse=True):
        if full_name == origin or full_name.startswith(origin + "."):
            return origin
    return ""


def jedi_lookup_column(node: ast.AST) -> int:
    if isinstance(node, ast.Attribute):
        return max(getattr(node, "end_col_offset", node.col_offset + len(node.attr)) - 1, 0)
    if isinstance(node, ast.Name):
        return node.col_offset
    return getattr(node, "col_offset", 0)


def jedi_sys_paths(repo: Path) -> list[str]:
    """Expose common source roots to Jedi without changing call analysis logic."""
    paths = [str(repo)]
    for child in ("src", "lib"):
        candidate = repo / child
        if candidate.exists():
            paths.append(str(candidate))
    return paths


def propagate_dynamic_loaders(
    files: dict[str, FileInfo],
    rules: list[LoaderRule],
    graph: nx.DiGraph,
) -> dict[str, DynamicLoaderResolution]:
    loader_strings: dict[str, LoaderStringResolution] = {}
    dynamic_loaders: dict[str, DynamicLoaderResolution] = {}
    functions = function_index(files)

    changed = True
    for _ in range(8):
        if not changed:
            break
        changed = False

        for file in files.values():
            for function in file.functions:
                for param, default in function.default_values.items():
                    resolution = loader_string_from_value(default, rules)
                    if not resolution:
                        continue
                    for symbol in parameter_symbols(file.module, function.class_name, function.function_name, param):
                        if symbol not in loader_strings:
                            loader_strings[symbol] = resolution
                            changed = True
                            add_dynamic_loader_string_node(graph, symbol, resolution, function.line_number)

            for assignment in file.assignments:
                resolution = loader_string_from_value(assignment.value, rules)
                if resolution:
                    for symbol in assignment_symbols(assignment, file.module):
                        if symbol not in loader_strings:
                            loader_strings[symbol] = resolution
                            changed = True
                            add_dynamic_loader_string_node(graph, symbol, resolution, assignment.line_number)

            for call in file.calls:
                target_function = resolve_called_function(call, file, functions)
                if not target_function:
                    continue
                for param, arg in call_argument_bindings(call, target_function):
                    resolution = expr_loader_string_resolution(
                        arg,
                        file,
                        call.class_name,
                        call.function_name,
                        loader_strings,
                    )
                    if not resolution:
                        continue
                    for symbol in parameter_symbols(
                        file.module,
                        target_function.class_name,
                        target_function.function_name,
                        param,
                    ):
                        if symbol not in loader_strings:
                            loader_strings[symbol] = resolution
                            changed = True
                            add_dynamic_loader_string_node(graph, symbol, resolution, call.node.lineno)

            for assignment in file.assignments:
                resolution = dynamic_loader_from_value(
                    assignment.value,
                    file,
                    assignment.class_name,
                    assignment.function_name,
                    loader_strings,
                )
                if not resolution:
                    continue
                for symbol in assignment_symbols(assignment, file.module):
                    if symbol not in dynamic_loaders:
                        dynamic_loaders[symbol] = resolution
                        changed = True
                        add_dynamic_loader_symbol_node(graph, symbol, resolution, assignment.line_number)

    return dynamic_loaders


def function_index(files: dict[str, FileInfo]) -> dict[str, FunctionDefRecord]:
    functions = {}
    for file in files.values():
        for function in file.functions:
            functions[function.symbol] = function
    return functions


def function_params(node: ast.FunctionDef) -> list[str]:
    args = list(node.args.posonlyargs) + list(node.args.args) + list(node.args.kwonlyargs)
    return [arg.arg for arg in args]


def function_default_values(node: ast.FunctionDef) -> dict[str, ast.AST]:
    values: dict[str, ast.AST] = {}
    positional = list(node.args.posonlyargs) + list(node.args.args)
    padded_defaults = [None] * (len(positional) - len(node.args.defaults)) + list(node.args.defaults)
    for arg, default in zip(positional, padded_defaults):
        if default is not None:
            values[arg.arg] = default
    for arg, default in zip(node.args.kwonlyargs, node.args.kw_defaults):
        if default is not None:
            values[arg.arg] = default
    return values


def is_property_like_function(node: ast.FunctionDef | ast.AsyncFunctionDef) -> bool:
    """Recognize descriptors that expose method returns as attributes."""
    for decorator in node.decorator_list:
        target = decorator.func if isinstance(decorator, ast.Call) else decorator
        if ".".join(expr_chain(target)) in {"property", "cached_property", "functools.cached_property"}:
            return True
    return False


def loader_string_from_value(node: ast.AST, rules: list[LoaderRule]) -> LoaderStringResolution | None:
    if not isinstance(node, ast.Constant) or not isinstance(node.value, str):
        return None
    matched_rules = rules_for_loader_string(node.value, rules)
    if not matched_rules:
        return None
    return LoaderStringResolution(
        origin=matched_rules[0].import_origin,
        loader_string=node.value,
        evidence=node.value,
        rules=tuple(matched_rules),
    )


def rules_for_loader_string(value: str, rules: list[LoaderRule]) -> list[LoaderRule]:
    normalized = value.strip()
    matches = []
    for rule in rules:
        patterns = loader_string_patterns(rule)
        if any(loader_string_matches(normalized, pattern) for pattern in patterns):
            matches.append(rule)
    return unique_rules(matches)


def loader_string_patterns(rule: LoaderRule) -> list[str]:
    patterns = []
    if rule.import_origin and rule.model_loader:
        patterns.append(f"{rule.import_origin}.{rule.model_loader}")
    for value in (rule.model_loader, rule.chain_suffix):
        if value and "." in value:
            patterns.append(value)
    return unique(patterns)


def loader_string_matches(value: str, pattern: str) -> bool:
    if not value or not pattern:
        return False
    return value == pattern or value.endswith("." + pattern)


def parameter_symbols(module: str, class_name: str, function_name: str, param: str) -> list[str]:
    symbols = [qualified_context_symbol(module, class_name, function_name, param)]
    if class_name:
        symbols.append(qualified_symbol(module, [class_name], param))
    else:
        symbols.append(qualified_symbol(module, [], param))
    return unique(symbols)


def assignment_symbols(assignment: AssignmentRecord, module: str) -> list[str]:
    symbols = [assignment.symbol]
    if assignment.target_name:
        symbols.append(
            qualified_context_symbol(
                module,
                assignment.class_name,
                assignment.function_name,
                assignment.target_name,
            )
        )
    return unique(symbols)


def qualified_context_symbol(module: str, class_name: str, function_name: str, name: str) -> str:
    parts = [module]
    if class_name:
        parts.append(class_name)
    if function_name:
        parts.append(function_name)
    parts.append(name)
    return ".".join(part for part in parts if part)


def expr_loader_string_resolution(
    node: ast.AST,
    file: FileInfo,
    class_name: str,
    function_name: str,
    loader_strings: dict[str, LoaderStringResolution],
) -> LoaderStringResolution | None:
    chain = expr_chain(node)
    if not chain:
        return None
    candidates = []
    if len(chain) == 1:
        candidates.append(qualified_context_symbol(file.module, class_name, function_name, chain[0]))
        if class_name:
            candidates.append(qualified_symbol(file.module, [class_name], chain[0]))
        candidates.append(qualified_symbol(file.module, [], chain[0]))
    for length in range(len(chain), 0, -1):
        candidates.append(qualified_symbol(file.module, [], ".".join(chain[:length])))
        if class_name:
            candidates.append(qualified_symbol(file.module, [class_name], ".".join(chain[:length])))
    for symbol in unique(candidates):
        if symbol in loader_strings:
            return loader_strings[symbol]
    return None


def dynamic_loader_from_value(
    node: ast.AST,
    file: FileInfo,
    class_name: str,
    function_name: str,
    loader_strings: dict[str, LoaderStringResolution],
) -> DynamicLoaderResolution | None:
    if not isinstance(node, ast.Call):
        return None
    func_chain = expr_chain(node.func)
    if not is_dynamic_loader_resolver(func_chain) or not node.args:
        return None
    string_resolution = expr_loader_string_resolution(
        node.args[0],
        file,
        class_name,
        function_name,
        loader_strings,
    )
    if not string_resolution:
        return None
    return DynamicLoaderResolution(
        origin=string_resolution.origin,
        method="dynamic_loader_string_resolver",
        evidence=f"{'.'.join(func_chain)}({string_resolution.evidence})",
        rules=string_resolution.rules,
    )


def is_dynamic_loader_resolver(func_chain: list[str]) -> bool:
    if not func_chain:
        return False
    chain = ".".join(func_chain)
    return chain in {"locate", "pydoc.locate", "import_string"} or chain.endswith(".import_string")


def resolve_called_function(
    call: CallRecord,
    file: FileInfo,
    functions: dict[str, FunctionDefRecord],
) -> FunctionDefRecord | None:
    parts = call.chain_parts
    if not parts:
        return None
    candidates = []
    if parts[0] == "self" and call.class_name and len(parts) >= 2:
        candidates.append(qualified_symbol(file.module, [call.class_name], parts[-1]))
    if len(parts) == 1:
        candidates.append(qualified_symbol(file.module, [], parts[0]))
        if call.class_name:
            candidates.append(qualified_symbol(file.module, [call.class_name], parts[0]))
    for symbol in unique(candidates):
        if symbol in functions:
            return functions[symbol]
    return None


def call_argument_bindings(call: CallRecord, function: FunctionDefRecord) -> list[tuple[str, ast.AST]]:
    params = list(function.params)
    if params and params[0] in {"self", "cls"}:
        params = params[1:]
    bindings: list[tuple[str, ast.AST]] = []
    for param, arg in zip(params, call.node.args):
        bindings.append((param, arg))
    for keyword in call.node.keywords:
        if keyword.arg and keyword.arg in params:
            bindings.append((keyword.arg, keyword.value))
    return bindings


def add_dynamic_loader_string_node(
    graph: nx.DiGraph,
    symbol: str,
    resolution: LoaderStringResolution,
    line_number: int,
) -> None:
    symbol_node = f"loader_string:{symbol}"
    graph.add_node(symbol_node, kind="loader_string", label=symbol)
    graph.add_edge(
        symbol_node,
        f"origin:{resolution.origin}",
        kind="loader_string_matches_origin",
        line=line_number,
        loader_string=resolution.loader_string,
    )


def add_dynamic_loader_symbol_node(
    graph: nx.DiGraph,
    symbol: str,
    resolution: DynamicLoaderResolution,
    line_number: int,
) -> None:
    symbol_node = f"dynamic_loader:{symbol}"
    graph.add_node(symbol_node, kind="dynamic_loader", label=symbol)
    graph.add_edge(
        f"origin:{resolution.origin}",
        symbol_node,
        kind=f"dynamic_loader_origin:{resolution.method}",
        line=line_number,
    )


def file_env(file: FileInfo, origins: dict[str, str]) -> dict[str, str]:
    env = {}
    for binding in file.imports.values():
        local_symbol = qualified_symbol(file.module, [], binding.alias)
        if binding.import_origin:
            env[binding.alias] = binding.import_origin
        elif binding.project_module:
            if binding.project_symbol:
                target = qualified_symbol(binding.project_module, [], binding.project_symbol)
                if target in origins:
                    env[binding.alias] = origins[target]
            else:
                for symbol, origin in origins.items():
                    if symbol.startswith(binding.project_module + "."):
                        env[binding.alias] = origin
                        break
        elif local_symbol in origins:
            env[binding.alias] = origins[local_symbol]

    prefix = file.module + "."
    for symbol, origin in origins.items():
        if symbol.startswith(prefix):
            local = symbol[len(prefix):]
            if "." not in local:
                env[local] = origin
    return env


def expr_origin_resolution(
    node: ast.AST,
    file: FileInfo,
    env: dict[str, str],
    origins: dict[str, str],
    origin_resolutions: dict[str, OriginResolution],
    class_name: str = "",
    function_name: str = "",
) -> OriginResolution:
    if isinstance(node, ast.BoolOp):
        for value in node.values:
            resolution = expr_origin_resolution(
                value,
                file,
                env,
                origins,
                origin_resolutions,
                class_name,
                function_name,
            )
            if resolution.origin:
                return OriginResolution(
                    resolution.origin,
                    f"boolop_fallback_origin:{resolution.method}",
                    resolution.evidence,
                )
    chain = expr_chain(node)
    if chain:
        if len(chain) == 1:
            context_symbol = qualified_context_symbol(file.module, class_name, function_name, chain[0])
            if context_symbol in origins:
                return origin_resolutions.get(
                    context_symbol,
                    OriginResolution(origins[context_symbol], "context_symbol_origin", chain[0]),
                )
            if class_name:
                class_symbol = qualified_symbol(file.module, [class_name], chain[0])
                if class_symbol in origins:
                    return origin_resolutions.get(
                        class_symbol,
                        OriginResolution(origins[class_symbol], "class_scoped_symbol_origin", chain[0]),
                    )
        if chain[0] == "self" and len(chain) >= 2:
            if class_name:
                symbol = f"{file.module}.{class_name}.{'.'.join(chain[1:])}"
                if symbol in origins:
                    return origin_resolutions.get(
                        symbol,
                        OriginResolution(origins[symbol], "class_field_origin", ".".join(chain)),
                    )
        for length in range(len(chain), 0, -1):
            symbol = qualified_symbol(file.module, [], ".".join(chain[:length]))
            if symbol in origins:
                return origin_resolutions.get(
                    symbol,
                    OriginResolution(origins[symbol], "local_symbol_origin", ".".join(chain[:length])),
                )
        if chain[0] in env:
            return OriginResolution(env[chain[0]], "import_or_project_alias", chain[0])
    if isinstance(node, ast.Call):
        func_chain = expr_chain(node.func)
        if func_chain:
            if func_chain[0] in env:
                return OriginResolution(env[func_chain[0]], "constructor_or_imported_call", ".".join(func_chain))
            if func_chain[0] == "self" and class_name and len(func_chain) >= 2:
                method_symbol = qualified_symbol(file.module, [class_name], func_chain[1], "return")
                if method_symbol in origins:
                    return origin_resolutions.get(
                        method_symbol,
                        OriginResolution(origins[method_symbol], "self_method_return_origin", ".".join(func_chain)),
                    )
            if class_name and len(func_chain) == 1:
                method_symbol = qualified_symbol(file.module, [class_name], func_chain[0], "return")
                if method_symbol in origins:
                    return origin_resolutions.get(
                        method_symbol,
                        OriginResolution(origins[method_symbol], "class_method_return_origin", ".".join(func_chain)),
                    )
            function_symbol = qualified_symbol(file.module, [], func_chain[0], "return")
            if function_symbol in origins:
                return origin_resolutions.get(
                    function_symbol,
                    OriginResolution(origins[function_symbol], "factory_return_origin", func_chain[0]),
                )
            if is_wrapper_like_call(func_chain):
                for child in list(node.args) + [keyword.value for keyword in node.keywords]:
                    child_resolution = expr_origin_resolution(
                        child,
                        file,
                        env,
                        origins,
                        origin_resolutions,
                        class_name,
                        function_name,
                    )
                    if child_resolution.origin:
                        return OriginResolution(
                            child_resolution.origin,
                            "wrapper_call_argument_origin",
                            f"{'.'.join(func_chain)}({child_resolution.evidence or child_resolution.origin})",
                        )
    return OriginResolution()


def is_wrapper_like_call(func_chain: list[str]) -> bool:
    if not func_chain:
        return False
    name = func_chain[-1].lower()
    wrapper_names = {
        "adapt",
        "adapter",
        "build",
        "configure",
        "factory",
        "init",
        "initialize",
        "make",
        "patch",
        "wrap",
        "wrapper",
    }
    wrapper_tokens = {
        "adapter",
        "client",
        "factory",
        "llm",
        "model",
        "wrapper",
    }
    return (
        name in wrapper_names
        or name.startswith(("from_", "to_", "as_", "with_"))
        or any(token in name for token in wrapper_tokens)
    )


def find_loader_candidates(
    files: dict[str, FileInfo],
    rules: list[LoaderRule],
    symbol_origins: dict[str, str],
    origin_resolutions: dict[str, OriginResolution],
    dynamic_loader_resolutions: dict[str, DynamicLoaderResolution],
    graph: nx.DiGraph,
    repo_full_name: str,
    commit: str,
) -> list[dict[str, object]]:
    suffix_rules: dict[str, list[LoaderRule]] = {}
    terminal_rules: dict[str, list[LoaderRule]] = {}
    all_model_args = set(DEFAULT_MODEL_ARGS)
    for rule in rules:
        if rule.chain_suffix:
            suffix_rules.setdefault(rule.chain_suffix, []).append(rule)
        if rule.terminal_call:
            terminal_rules.setdefault(rule.terminal_call, []).append(rule)
        all_model_args.update(rule.model_args)

    candidates = []
    candidate_index = 0
    for file in files.values():
        env = file_env(file, symbol_origins)
        for call in file.calls:
            terminal = call.chain_parts[-1]
            matched_suffix_rules = rules_for_chain(call.visible_chain, suffix_rules)
            matched_terminal_rules = terminal_rules.get(terminal, [])
            matched_rules = unique_rules(matched_suffix_rules or matched_terminal_rules)
            has_model_arg = bool(set(call.keyword_args) & all_model_args) or call_has_positional_model_arg(
                call,
                matched_rules,
            )
            has_model_payload = call.has_model_payload or call_uses_model_payload_symbol(call, file)
            receiver_symbol = receiver_for_call(call)
            receiver_resolution = resolve_receiver_origin(call, file, env, symbol_origins, origin_resolutions)
            dynamic_resolution = resolve_dynamic_loader_for_call(call, file, dynamic_loader_resolutions)
            if dynamic_resolution:
                matched_rules = unique_rules(dynamic_resolution.rules)
                matched_suffix_rules = matched_rules
                receiver_symbol = receiver_symbol or call.visible_chain
                receiver_resolution = OriginResolution(
                    origin=dynamic_resolution.origin,
                    method=dynamic_resolution.method,
                    evidence=dynamic_resolution.evidence,
                )
                receiver_origin = receiver_resolution.origin
            elif not matched_rules:
                continue
            receiver_origin = receiver_resolution.origin
            if receiver_origin:
                matched_suffix_rules = origin_compatible_rules(matched_suffix_rules, receiver_origin)
                matched_terminal_rules = origin_compatible_rules(matched_terminal_rules, receiver_origin)
                matched_rules = unique_rules(matched_suffix_rules or matched_terminal_rules)
            else:
                matched_suffix_rules = []
                matched_rules = []
            if not matched_rules:
                continue
            if dynamic_resolution:
                binding_eligible = True
            else:
                binding_eligible = classify_call_candidate(
                    call=call,
                    matched_suffix_rules=matched_suffix_rules,
                    matched_rules=matched_rules,
                    receiver_origin=receiver_origin,
                    has_model_arg=has_model_arg,
                    has_model_payload=has_model_payload,
                )
            if binding_eligible is None:
                continue
            matched_rule_origins = unique(rule.import_origin for rule in matched_rules)
            candidate_index += 1
            candidate_id = f"loader_candidate_{candidate_index:05d}"
            matched_loader_values = unique(rule.model_loader for rule in matched_rules)
            matched_suffixes = unique(rule.chain_suffix for rule in matched_rules)
            location_url = make_location_url(repo_full_name, commit, call.file_path, call.node.lineno)
            row = {
                "loader_candidate_id": candidate_id,
                "file_path": call.file_path,
                "line_number": call.node.lineno,
                "location_url": location_url,
                "visible_call_chain": call.visible_chain,
                "terminal_call": terminal,
                "matched_chain_suffix": "|".join(matched_suffixes),
                "matched_canonical_loader": "|".join(matched_loader_values),
                "matched_rule_import_origin": "|".join(matched_rule_origins),
                "linked_import_origin": receiver_origin,
                "receiver_symbol": receiver_symbol,
                "receiver_origin": receiver_origin,
                "origin_resolution_method": receiver_resolution.method,
                "origin_resolution_evidence": receiver_resolution.evidence,
                "has_model_arg": has_model_arg,
                "has_model_payload": has_model_payload,
                "binding_eligible": binding_eligible,
                "line_text": call.line_text,
            }
            candidates.append(row)
            add_call_graph_nodes(graph, candidate_id, call, receiver_origin, matched_loader_values, binding_eligible)
    return candidates


def origin_compatible_rules(rules: list[LoaderRule], receiver_origin: str) -> list[LoaderRule]:
    """Keep only codebook rules for the already-resolved receiver origin."""
    if not receiver_origin:
        return []
    return [rule for rule in rules if rule.import_origin == receiver_origin]


def classify_call_candidate(
    call: CallRecord,
    matched_suffix_rules: list[LoaderRule],
    matched_rules: list[LoaderRule],
    receiver_origin: str,
    has_model_arg: bool,
    has_model_payload: bool,
) -> bool | None:
    """Return sink eligibility, or None when the call is not a candidate."""
    terminal = call.chain_parts[-1]
    has_model_evidence = has_model_arg or has_model_payload
    matched_origins = {rule.import_origin for rule in matched_rules}
    receiver_matches_rule = bool(receiver_origin and receiver_origin in matched_origins)
    http_like_only = bool(matched_rules) and all(
        rule.import_origin in {"httpx", "requests", "urllib"} for rule in matched_rules
    )
    single_token_suffix_only = bool(matched_suffix_rules) and all(
        "." not in rule.chain_suffix for rule in matched_suffix_rules
    )

    if http_like_only and not has_model_payload:
        return None
    if single_token_suffix_only and not receiver_matches_rule:
        return None
    if not receiver_matches_rule:
        return None
    if matched_suffix_rules and receiver_matches_rule:
        return True
    if receiver_matches_rule and has_model_evidence and terminal not in GENERIC_TERMINALS:
        return True
    if terminal in GENERIC_TERMINALS and has_model_evidence and receiver_matches_rule:
        return False
    return None


def call_has_positional_model_arg(call: CallRecord, matched_rules: list[LoaderRule]) -> bool:
    if not call.node.args:
        return False
    terminal = call.chain_parts[-1] if call.chain_parts else ""
    normalized_terminal = terminal.lower()
    positional_model_selector = (
        terminal[:1].isupper()
        or normalized_terminal in {"get_model", "load_model", "from_pretrained"}
        or normalized_terminal.endswith("_model")
    )
    return positional_model_selector and any(rule.model_args for rule in matched_rules)


def rules_for_chain(visible_chain: str, suffix_rules: dict[str, list[LoaderRule]]) -> list[LoaderRule]:
    matches = []
    for suffix, rules in suffix_rules.items():
        if visible_chain == suffix or visible_chain.endswith("." + suffix):
            matches.extend(rules)
    return unique_rules(matches)


def receiver_for_call(call: CallRecord) -> str:
    if len(call.chain_parts) <= 1:
        return ""
    return ".".join(call.chain_parts[:-1])


def resolve_receiver_origin(
    call: CallRecord,
    file: FileInfo,
    env: dict[str, str],
    origins: dict[str, str],
    origin_resolutions: dict[str, OriginResolution],
) -> OriginResolution:
    parts = call.chain_parts
    if not parts:
        return OriginResolution()
    if parts[0] == "self" and call.class_name and len(parts) >= 2:
        for length in range(len(parts) - 1, 1, -1):
            symbol = f"{file.module}.{call.class_name}.{'.'.join(parts[1:length])}"
            if symbol in origins:
                return origin_resolutions.get(
                    symbol,
                    OriginResolution(origins[symbol], "class_field_origin", ".".join(parts[:length])),
                )
    for length in range(len(parts) - 1, 0, -1):
        symbol = qualified_symbol(file.module, [], ".".join(parts[:length]))
        if symbol in origins:
            return origin_resolutions.get(
                symbol,
                OriginResolution(origins[symbol], "local_symbol_origin", ".".join(parts[:length])),
            )
        if call.class_name:
            class_symbol = qualified_symbol(file.module, [call.class_name], ".".join(parts[:length]))
            if class_symbol in origins:
                return origin_resolutions.get(
                    class_symbol,
                    OriginResolution(origins[class_symbol], "class_scoped_local_origin", ".".join(parts[:length])),
                )
    if parts[0] in env:
        return OriginResolution(env[parts[0]], "import_or_project_alias", parts[0])
    return OriginResolution()


def resolve_dynamic_loader_for_call(
    call: CallRecord,
    file: FileInfo,
    dynamic_loaders: dict[str, DynamicLoaderResolution],
) -> DynamicLoaderResolution | None:
    parts = call.chain_parts
    if not parts:
        return None
    candidates = []
    if len(parts) == 1:
        candidates.append(qualified_context_symbol(file.module, call.class_name, call.function_name, parts[0]))
        if call.class_name:
            candidates.append(qualified_symbol(file.module, [call.class_name], parts[0]))
        candidates.append(qualified_symbol(file.module, [], parts[0]))
    for length in range(len(parts), 0, -1):
        fragment = ".".join(parts[:length])
        candidates.append(qualified_context_symbol(file.module, call.class_name, call.function_name, fragment))
        if call.class_name:
            candidates.append(qualified_symbol(file.module, [call.class_name], fragment))
        candidates.append(qualified_symbol(file.module, [], fragment))
    for symbol in unique(candidates):
        if symbol in dynamic_loaders:
            return dynamic_loaders[symbol]
    return None


def add_origin_edge(graph: nx.DiGraph, origin: str, symbol: str, edge_kind: str, line_number: int) -> None:
    origin_node = f"origin:{origin}"
    symbol_node = f"symbol:{symbol}"
    graph.add_node(origin_node, kind="import_origin", label=origin)
    graph.add_node(symbol_node, kind="symbol", label=symbol)
    graph.add_edge(origin_node, symbol_node, kind=edge_kind, line=line_number)


def add_call_graph_nodes(
    graph: nx.DiGraph,
    candidate_id: str,
    call: CallRecord,
    receiver_origin: str,
    matched_loaders: list[str],
    binding_eligible: bool,
) -> None:
    call_node = f"call:{call.file_path}:{call.node.lineno}:{call.node.col_offset}"
    candidate_node = f"loader_candidate:{candidate_id}"
    graph.add_node(call_node, kind="call", label=call.visible_chain)
    graph.add_node(candidate_node, kind="loader_candidate", label=candidate_id, binding_eligible=binding_eligible)
    graph.add_edge(call_node, candidate_node, kind="is_loader_candidate")
    for loader in matched_loaders:
        loader_node = f"loader:{loader}"
        graph.add_node(loader_node, kind="loader_signature", label=loader)
        graph.add_edge(candidate_node, loader_node, kind="matches_loader")
    if receiver_origin:
        graph.add_edge(f"origin:{receiver_origin}", call_node, kind="origin_reaches_call")


def write_graph_edges(path: Path, graph: nx.DiGraph) -> None:
    rows = []
    for source, target, data in sorted(graph.edges(data=True), key=lambda edge: (edge[0], edge[1], str(edge[2]))):
        rows.append(
            {
                "source": source,
                "target": target,
                "edge_kind": data.get("kind", ""),
                "line": data.get("line", ""),
            }
        )
    write_csv(path, rows)


def choose_module_name(path: Path, path_to_modules: dict[Path, list[str]]) -> str:
    modules = sorted(path_to_modules.get(path, []), key=lambda value: (value.startswith("src."), len(value)))
    return modules[0] if modules else path.stem


def resolve_relative_module(current_module: str, module: str, level: int) -> str:
    if level <= 0:
        return module
    parts = current_module.split(".")
    base = parts[: max(len(parts) - level, 0)]
    if module:
        base.extend(module.split("."))
    return ".".join(part for part in base if part)


def qualified_symbol(module: str, class_stack: list[str], name: str, suffix: str = "") -> str:
    parts = [module]
    parts.extend(class_stack)
    parts.append(name)
    if suffix:
        parts.append(suffix)
    return ".".join(part for part in parts if part)


def class_name_from_symbol_context(file: FileInfo, node: ast.AST) -> str:
    # Receiver resolution uses CallRecord.class_name. An isolated expression
    # does not establish its enclosing class here.
    return ""


def expr_chain(node: ast.AST) -> list[str]:
    if isinstance(node, ast.Name):
        return [node.id]
    if isinstance(node, ast.Attribute):
        return expr_chain(node.value) + [node.attr]
    if isinstance(node, ast.Call):
        return expr_chain(node.func)
    if isinstance(node, ast.Subscript):
        return expr_chain(node.value)
    return []


def call_has_model_payload(node: ast.Call) -> bool:
    for keyword in node.keywords:
        if keyword.arg in {"json", "data", "body"} and ast_contains_model_key(keyword.value):
            return True
    return False


def call_uses_model_payload_symbol(call: CallRecord, file: FileInfo) -> bool:
    """Detect request bodies built as variables before the loader call."""
    for keyword in call.node.keywords:
        if keyword.arg not in {"json", "data", "body"}:
            continue
        if expr_uses_model_payload_symbol(keyword.value, call, file):
            return True
    return False


def expr_uses_model_payload_symbol(expr: ast.AST, call: CallRecord, file: FileInfo) -> bool:
    if isinstance(expr, ast.Name):
        return bool(candidate_value_symbols(file, call, expr.id) & file.model_payload_symbols)
    if isinstance(expr, ast.Attribute):
        chain = expr_chain(expr)
        if chain and chain[0] == "self" and call.class_name:
            symbol = qualified_symbol(file.module, [call.class_name], ".".join(chain[1:]))
            return symbol in file.model_payload_symbols
        symbol = qualified_symbol(file.module, [], ".".join(chain))
        return symbol in file.model_payload_symbols
    if isinstance(expr, ast.Dict):
        return ast_contains_model_key(expr) or any(
            expr_uses_model_payload_symbol(value, call, file)
            for value in expr.values
        )
    if isinstance(expr, (ast.List, ast.Tuple, ast.Set)):
        return any(expr_uses_model_payload_symbol(item, call, file) for item in expr.elts)
    return False


def candidate_value_symbols(file: FileInfo, call: CallRecord, name: str) -> set[str]:
    symbols = {qualified_symbol(file.module, [], name)}
    if call.class_name:
        symbols.add(qualified_symbol(file.module, [call.class_name], name))
    return symbols


def ast_contains_model_key(node: ast.AST) -> bool:
    """Return whether an AST subtree contains a literal ``model`` dict key.

    Some generated or machine-written files contain extremely deep expression
    trees, so this deliberately uses an explicit stack instead of recursion.
    """
    stack = [node]
    seen: set[int] = set()
    while stack:
        current = stack.pop()
        object_id = id(current)
        if object_id in seen:
            continue
        seen.add(object_id)
        if isinstance(current, ast.Dict):
            for key in current.keys:
                if isinstance(key, ast.Constant) and str(key.value) == "model":
                    return True
        stack.extend(ast.iter_child_nodes(current))
    return False


def line_at(lines: list[str], line_number: int) -> str:
    if 1 <= line_number <= len(lines):
        return lines[line_number - 1].strip()
    return ""


def make_location_url(repo_full_name: str, commit: str, file_path: str, line_number: int) -> str:
    if not repo_full_name or not commit:
        return ""
    return f"https://github.com/{repo_full_name}/blob/{commit}/{file_path}#L{line_number}"


def unique(values: Iterable[str]) -> list[str]:
    seen = set()
    output = []
    for value in values:
        if not value or value in seen:
            continue
        seen.add(value)
        output.append(value)
    return output


def unique_rules(rules: Iterable[LoaderRule]) -> list[LoaderRule]:
    seen = set()
    output = []
    for rule in rules:
        key = (rule.import_origin, rule.model_loader)
        if key in seen:
            continue
        seen.add(key)
        output.append(rule)
    return output


def cell_text(value: object) -> str:
    return str(value or "").strip()


if __name__ == "__main__":
    raise SystemExit(main())
