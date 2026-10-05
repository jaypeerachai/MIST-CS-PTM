#!/usr/bin/env python
"""Find exact PTM IDs and import origins in repository source."""

from __future__ import annotations

from mist.resources import DATA_DIR

import argparse
import ast
import csv
import hashlib
import io
import json
import re
import tokenize
from dataclasses import asdict, dataclass, fields
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook

from mist.io import display_output_path, read_csv, write_csv
from mist.analysis.parse_cache import parse_python_files
from mist.rules.path_terms import EXAMPLE_DEMO_TOKENS, FILENAME_EXAMPLE_DEMO_TOKENS, THIRD_PARTY_TOKENS


@dataclass(frozen=True)
class ModelIdOccurrence:
    canonical_model_id: str
    matched_text: str
    variant: str
    file_path: str
    line_number: int
    column_start: int
    column_end: int
    line_text: str
    path_signal: str
    source_context: str
    fp_signal: str
    excluded_from_main_queue: bool
    occurrence_decision: str


@dataclass(frozen=True)
class ImportOriginOccurrence:
    import_origin: str
    import_style: str
    imported_name: str
    alias: str
    file_path: str
    line_number: int
    line_text: str


def main(argv: Sequence[str] | None = None) -> int:
    """Scan one repo snapshot and write the identifier extraction seed-node CSVs."""
    args = parse_args(argv)
    repo = args.repo.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    model_ids = load_model_ids(args.model_ids)
    import_origins = load_import_origins(args.reuse_codebook)
    model_terms = build_model_terms(model_ids)

    model_rows: list[ModelIdOccurrence] = []
    import_rows: list[ImportOriginOccurrence] = []
    if args.syntax_cache:
        python_files, parse_errors, model_rows, import_rows = scan_from_syntax_cache(
            repo,
            args.syntax_cache.resolve(),
            model_terms,
            import_origins,
        )
    else:
        python_files, parse_errors, model_rows, import_rows = scan_live_repo(
            repo,
            model_terms,
            import_origins,
        )

    for name, rows, record_type in (
        ("model_id_occurrences.csv", model_rows, ModelIdOccurrence),
        ("import_origin_occurrences.csv", import_rows, ImportOriginOccurrence),
    ):
        with (output_dir / name).open("w", newline="", encoding="utf-8", errors="backslashreplace") as handle:
            writer = csv.DictWriter(handle, fieldnames=[field.name for field in fields(record_type)],
                                    escapechar="\\")
            writer.writeheader()
            writer.writerows(asdict(row) for row in rows)
    write_csv(
        output_dir / "occurrence_seed_nodes.csv",
        normalize_combined_rows(model_rows, import_rows),
    )

    summary = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "repo": str(repo),
        "model_ids_file": str(args.model_ids),
        "reuse_codebook": str(args.reuse_codebook),
        "syntax_cache": str(args.syntax_cache) if args.syntax_cache else "",
        "python_files_scanned": python_files,
        "python_parse_errors": parse_errors,
        "model_ids_loaded": len(model_ids),
        "model_search_terms": len(model_terms),
        "import_origins_loaded": len(import_origins),
        "model_id_occurrences": len(model_rows),
        "model_id_occurrences_excluded_from_main_queue": sum(
            row.excluded_from_main_queue for row in model_rows
        ),
        "model_id_occurrences_docstring_comment": sum(
            row.fp_signal == "docstring_comment" for row in model_rows
        ),
        "import_origin_occurrences": len(import_rows),
        "outputs": {
            "model_id_occurrences": display_output_path(output_dir / "model_id_occurrences.csv", resolve=True),
            "import_origin_occurrences": display_output_path(output_dir / "import_origin_occurrences.csv", resolve=True),
            "occurrence_seed_nodes": display_output_path(output_dir / "occurrence_seed_nodes.csv", resolve=True),
            "summary": display_output_path(output_dir / "occurrence_summary.json", resolve=True),
        },
    }
    (output_dir / "occurrence_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(json.dumps(summary, indent=2))
    return 0


def parse_args(argv: Sequence[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract occurrence discovery model-ID and import-origin seed nodes.")
    parser.add_argument("--repo", type=Path, required=True, help="Repository snapshot root.")
    parser.add_argument("--model-ids", type=Path, required=True, help="CSV file with a model_id column.")
    parser.add_argument(
        "--reuse-codebook",
        type=Path,
        default=DATA_DIR / "reuse_rules.xlsx",
        help="XLSX codebook with import_origin_seeds sheet.",
    )
    parser.add_argument("--output-dir", type=Path, required=True, help="Output directory.")
    parser.add_argument(
        "--syntax-cache",
        type=Path,
        help="Optional syntax cache directory with files/imports/source_context_spans CSVs.",
    )
    return parser.parse_args(argv)


def load_model_ids(path: Path) -> list[str]:
    seen = set()
    ids = []
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if "model_id" not in (reader.fieldnames or []):
            raise ValueError(f"{path} must have a model_id column")
        for row in reader:
            model_id = (row.get("model_id") or "").strip()
            if model_id and model_id not in seen:
                seen.add(model_id)
                ids.append(model_id)
    return ids


def load_import_origins(path: Path) -> set[str]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["import_origin_seeds"]
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    origin_col = headers.index("import origin") + 1
    origins = set()
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, origin_col).value
        if value:
            origins.add(str(value).strip())
    return origins


def build_model_terms(model_ids: Iterable[str]) -> list[tuple[str, str, str]]:
    terms = []
    seen = set()
    for model_id in model_ids:
        variants = [
            (f"'{model_id}'", "full_single_quoted"),
            (f'"{model_id}"', "full_double_quoted"),
        ]
        if "/" in model_id:
            bare = model_id.rsplit("/", 1)[-1]
            variants.extend(
                [
                    (f"'{bare}'", "namespace_free_single_quoted"),
                    (f'"{bare}"', "namespace_free_double_quoted"),
                ]
            )
        for needle, variant in variants:
            key = (model_id, needle, variant)
            if key in seen:
                continue
            seen.add(key)
            terms.append((model_id, needle, variant))
    return terms


def scan_live_repo(
    repo: Path,
    model_terms: list[tuple[str, str, str]],
    import_origins: set[str],
) -> tuple[int, int, list[ModelIdOccurrence], list[ImportOriginOccurrence]]:
    """Scan source files directly, using AST only for source spans and imports."""
    model_rows: list[ModelIdOccurrence] = []
    import_rows: list[ImportOriginOccurrence] = []
    python_files = 0
    parse_errors = 0

    for parsed in parse_python_files(repo):
        python_files += 1
        rel_path = parsed.rel_path
        path_signal = detect_path_signal(rel_path) or ""
        text = parsed.text
        lines = parsed.lines
        tree = parsed.tree

        if parsed.parse_failed:
            parse_errors += 1
        # Keep comments/docstrings for audit, but mark them so binding tracing can skip them.
        source_spans = collect_source_context_spans(text, tree)
        model_rows.extend(scan_model_ids(rel_path, lines, model_terms, path_signal, source_spans))
        if tree is None:
            continue
        # Imports are AST-based so aliases and from-imports are captured consistently.
        import_rows.extend(scan_import_origins(rel_path, lines, tree, import_origins))

    return python_files, parse_errors, model_rows, import_rows


def scan_from_syntax_cache(
    repo: Path,
    cache_dir: Path,
    model_terms: list[tuple[str, str, str]],
    import_origins: set[str],
) -> tuple[int, int, list[ModelIdOccurrence], list[ImportOriginOccurrence]]:
    """Reuse syntax index facts for spans/imports while keeping source text current."""
    file_rows = read_csv(cache_dir / "files.csv")
    spans_by_file = load_cached_source_spans(cache_dir / "source_context_spans.csv")
    imports_by_file = load_cached_imports(cache_dir / "imports.csv")

    model_rows: list[ModelIdOccurrence] = []
    import_rows: list[ImportOriginOccurrence] = []
    parse_errors = 0

    for row in file_rows:
        rel_path = row.get("file_path", "")
        if not rel_path:
            continue
        source_path = repo / rel_path
        text = source_path.read_text(encoding="utf-8", errors="replace")
        cached_sha = row.get("sha256", "")
        current_sha = hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()
        if cached_sha and cached_sha != current_sha:
            raise ValueError(f"syntax parsing cache is stale for {rel_path}")

        path_signal = detect_path_signal(rel_path) or ""
        lines = text.splitlines()
        model_rows.extend(
            scan_model_ids(
                rel_path,
                lines,
                model_terms,
                path_signal,
                spans_by_file.get(rel_path, []),
            )
        )
        import_rows.extend(
            import_origin_rows_from_cache(
                rel_path,
                imports_by_file.get(rel_path, []),
                import_origins,
            )
        )
        if row.get("parse_error", ""):
            parse_errors += 1

    return len(file_rows), parse_errors, model_rows, import_rows


def load_cached_source_spans(path: Path) -> dict[str, list[tuple[str, int, int, int, int]]]:
    spans_by_file: dict[str, list[tuple[str, int, int, int, int]]] = {}
    for row in read_csv(path):
        rel_path = row.get("file_path", "")
        if not rel_path:
            continue
        spans_by_file.setdefault(rel_path, []).append(
            (
                row.get("source_context", ""),
                int(row.get("start_line") or 0),
                int(row.get("start_col") or 0),
                int(row.get("end_line") or 0),
                int(row.get("end_col") or 0),
            )
        )
    return spans_by_file


def load_cached_imports(path: Path) -> dict[str, list[dict[str, str]]]:
    imports_by_file: dict[str, list[dict[str, str]]] = {}
    for row in read_csv(path):
        rel_path = row.get("file_path", "")
        if rel_path:
            imports_by_file.setdefault(rel_path, []).append(row)
    return imports_by_file


def import_origin_rows_from_cache(
    rel_path: str,
    cached_imports: list[dict[str, str]],
    import_origins: set[str],
) -> list[ImportOriginOccurrence]:
    rows: list[ImportOriginOccurrence] = []
    for cached in cached_imports:
        root = cached.get("root", "")
        if root not in import_origins:
            continue
        rows.append(
            ImportOriginOccurrence(
                import_origin=root,
                import_style=cached.get("import_style", ""),
                imported_name=cached.get("imported_name", ""),
                alias=cached.get("alias", ""),
                file_path=rel_path,
                line_number=int(cached.get("line_number") or 0),
                line_text=cached.get("line_text", ""),
            )
        )
    return rows


def scan_model_ids(
    rel_path: str,
    lines: list[str],
    model_terms: list[tuple[str, str, str]],
    path_signal: str,
    source_spans: list[tuple[str, int, int, int, int]],
) -> list[ModelIdOccurrence]:
    rows = []
    for line_number, line in enumerate(lines, start=1):
        for canonical_model_id, needle, variant in model_terms:
            start = 0
            while True:
                index = line.find(needle, start)
                if index == -1:
                    break
                end = index + len(needle)
                source_context = detect_source_context(source_spans, line_number, index, end)
                fp_signal = "docstring_comment" if source_context in {"comment", "docstring"} else ""
                excluded = (
                    path_signal in {"example_or_demo_code", "third_party_or_vendored_code"}
                    or fp_signal == "docstring_comment"
                )
                rows.append(
                    ModelIdOccurrence(
                        canonical_model_id=canonical_model_id,
                        matched_text=needle,
                        variant=variant,
                        file_path=rel_path,
                        line_number=line_number,
                        column_start=index + 1,
                        column_end=end,
                        line_text=line.strip(),
                        path_signal=path_signal,
                        source_context=source_context,
                        fp_signal=fp_signal,
                        excluded_from_main_queue=excluded,
                        occurrence_decision="exclude" if excluded else "include",
                    )
                )
                start = end
    return rows


def collect_source_context_spans(
    text: str,
    tree: ast.AST | None,
) -> list[tuple[str, int, int, int, int]]:
    spans: list[tuple[str, int, int, int, int]] = []
    try:
        tokens = tokenize.generate_tokens(io.StringIO(text).readline)
        for token in tokens:
            if token.type == tokenize.COMMENT:
                start_line, start_col = token.start
                end_line, end_col = token.end
                spans.append(("comment", start_line, start_col, end_line, end_col))
    except (tokenize.TokenError, IndentationError):
        pass

    if tree is not None:
        for node in ast.walk(tree):
            if isinstance(node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                doc_expr = docstring_expr(node)
                if doc_expr is None:
                    continue
                if not all(
                    hasattr(doc_expr, attr)
                    for attr in ("lineno", "col_offset", "end_lineno", "end_col_offset")
                ):
                    continue
                spans.append(
                    (
                        "docstring",
                        doc_expr.lineno,
                        doc_expr.col_offset,
                        doc_expr.end_lineno,
                        doc_expr.end_col_offset,
                    )
                )
    return spans


def docstring_expr(node: ast.AST) -> ast.Expr | None:
    body = getattr(node, "body", None)
    if not body:
        return None
    first = body[0]
    if not isinstance(first, ast.Expr):
        return None
    value = first.value
    if isinstance(value, ast.Constant) and isinstance(value.value, str):
        return first
    return None


def detect_source_context(
    spans: list[tuple[str, int, int, int, int]],
    line_number: int,
    start_col: int,
    end_col: int,
) -> str:
    for context, start_line, span_start_col, end_line, span_end_col in spans:
        if line_number < start_line or line_number > end_line:
            continue
        effective_start_col = span_start_col if line_number == start_line else 0
        effective_end_col = span_end_col if line_number == end_line else 10**9
        if start_col >= effective_start_col and end_col <= effective_end_col:
            return context
    return "code"


def scan_import_origins(
    rel_path: str,
    lines: list[str],
    tree: ast.AST,
    import_origins: set[str],
) -> list[ImportOriginOccurrence]:
    rows = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                root = alias.name.split(".", 1)[0]
                if root in import_origins:
                    rows.append(
                        ImportOriginOccurrence(
                            import_origin=root,
                            import_style="import",
                            imported_name=alias.name,
                            alias=alias.asname or "",
                            file_path=rel_path,
                            line_number=node.lineno,
                            line_text=line_at(lines, node.lineno),
                        )
                    )
        elif isinstance(node, ast.ImportFrom) and node.module:
            root = node.module.split(".", 1)[0]
            if root in import_origins:
                for alias in node.names:
                    rows.append(
                        ImportOriginOccurrence(
                            import_origin=root,
                            import_style="from",
                            imported_name=f"{node.module}.{alias.name}",
                            alias=alias.asname or "",
                            file_path=rel_path,
                            line_number=node.lineno,
                            line_text=line_at(lines, node.lineno),
                        )
                    )
    return rows


def line_at(lines: list[str], line_number: int) -> str:
    if 1 <= line_number <= len(lines):
        return lines[line_number - 1].strip()
    return ""


def detect_path_signal(rel_path: str) -> str | None:
    directory_tokens, filename_tokens = split_path_tokens(rel_path)
    if directory_tokens & EXAMPLE_DEMO_TOKENS:
        return "example_or_demo_code"
    if filename_tokens & FILENAME_EXAMPLE_DEMO_TOKENS:
        return "example_or_demo_code"
    tokens = directory_tokens | filename_tokens
    if tokens & THIRD_PARTY_TOKENS:
        return "third_party_or_vendored_code"
    return None


def split_path_tokens(rel_path: str) -> tuple[set[str], set[str]]:
    parts = Path(rel_path).parts
    if not parts:
        return set(), set()
    return path_part_tokens(parts[:-1]), path_part_tokens(parts[-1:])


def path_tokens(rel_path: str) -> set[str]:
    directory_tokens, filename_tokens = split_path_tokens(rel_path)
    return directory_tokens | filename_tokens


def path_part_tokens(parts: Iterable[str]) -> set[str]:
    tokens = set()
    for part in parts:
        lower = part.lower()
        tokens.add(lower)
        tokens.update(token for token in re.split(r"[_\-.]", lower) if token)
    return tokens


def normalize_combined_rows(
    model_rows: list[ModelIdOccurrence],
    import_rows: list[ImportOriginOccurrence],
) -> list[dict[str, object]]:
    rows = []
    for row in model_rows:
        data = asdict(row)
        rows.append(
            {
                "node_kind": "model_id_occurrence",
                "file_path": data.pop("file_path"),
                "line_number": data.pop("line_number"),
                "primary_value": data.pop("canonical_model_id"),
                "secondary_value": data.pop("matched_text"),
                "details_json": json.dumps(data, ensure_ascii=False),
            }
        )
    for row in import_rows:
        data = asdict(row)
        rows.append(
            {
                "node_kind": "import_origin_occurrence",
                "file_path": data.pop("file_path"),
                "line_number": data.pop("line_number"),
                "primary_value": data.pop("import_origin"),
                "secondary_value": data.pop("imported_name"),
                "details_json": json.dumps(data, ensure_ascii=False),
            }
        )
    return rows


if __name__ == "__main__":
    raise SystemExit(main())
