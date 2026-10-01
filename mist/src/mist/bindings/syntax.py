"""Syntax helpers for symbols, calls, configuration, and mock targets."""

from __future__ import annotations

from collections import defaultdict
from mist.io import display_output_path
from mist.io import parse_int
from mist.rules.binding_terms import DERIVED_TERM_STOPWORDS
from mist.rules.binding_terms import GENERIC_CONFIG_ARG_NAMES
from mist.rules.binding_terms import GENERIC_CONFIG_TERMS
from mist.rules.binding_terms import GENERIC_ENDPOINT_TERMS
from mist.rules.binding_terms import MOCK_SYMBOL_TERMS
from pathlib import Path
from typing import Any
from typing import Iterable
import ast
import re
from mist.bindings.constants import (
    MODEL_ARG_NAMES,
    MODEL_PAYLOAD_KEYS,
)
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from mist.bindings.graph import SourceSinkBuilder
from mist.bindings.types import (
    ClassDefInfo,
    FileFacts,
    IdentityVocabulary,
    MockPatchIndex,
    RepoFacts,
    SinkInfo,
    BindingLoaderRule,
)


def build_identity_vocabulary(codebook_path: Path, loader_rows: list[dict[str, str]]) -> IdentityVocabulary:
    provider_terms: set[str] = set()
    loader_action_terms: set[str] = set()
    loader_identity_terms: set[str] = set()
    codebook_rows_loaded = 0
    codebook_status = "not_found"

    if codebook_path.exists():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(codebook_path, read_only=True, data_only=True)
            for sheet in wb.worksheets:
                if sheet.title not in {"import_origin_seeds", "model_loading"}:
                    continue
                rows = sheet.iter_rows(values_only=True)
                try:
                    header = next(rows)
                except StopIteration:
                    continue
                header_map = {
                    normalize_header(value): index
                    for index, value in enumerate(header)
                    if value is not None
                }
                for values in rows:
                    row = {
                        name: clean_xlsx_value(values[index] if index < len(values) else "")
                        for name, index in header_map.items()
                    }
                    if not any(row.values()):
                        continue
                    add_provider_terms(provider_terms, row.get("import origin", ""))
                    add_loader_action_terms(loader_action_terms, row.get("terminal_call", ""))
                    add_loader_action_terms(loader_action_terms, row.get("chain_suffix", ""))
                    add_loader_identity_terms(loader_identity_terms, row.get("model loader", ""))
                    add_loader_identity_terms(loader_identity_terms, row.get("chain_suffix", ""))
                    codebook_rows_loaded += 1
            codebook_status = "loaded"
        except Exception as exc:
            codebook_status = f"load_error:{exc.__class__.__name__}:{exc}"

    for row in loader_rows:
        for field_name in ("linked_import_origin", "receiver_origin"):
            value = row.get(field_name, "")
            if value:
                add_provider_terms(provider_terms, value)
        for field_name in ("terminal_call", "matched_chain_suffix"):
            add_loader_action_terms(loader_action_terms, row.get(field_name, ""))
        add_loader_action_terms(loader_action_terms, terminal_from_chain(row.get("visible_call_chain", "")))
        for field_name in ("matched_canonical_loader", "visible_call_chain", "matched_chain_suffix"):
            add_loader_identity_terms(loader_identity_terms, row.get(field_name, ""))

    provider_terms = sanitize_derived_terms(provider_terms, keep_stopwords=False)
    loader_action_terms = sanitize_derived_terms(loader_action_terms, keep_stopwords=False)
    loader_identity_terms = sanitize_derived_terms(loader_identity_terms, keep_stopwords=False)
    endpoint_terms = sanitize_derived_terms(loader_action_terms | GENERIC_ENDPOINT_TERMS, keep_stopwords=True)

    return IdentityVocabulary(
        provider_terms=frozenset(provider_terms),
        loader_action_terms=frozenset(loader_action_terms),
        loader_identity_terms=frozenset(loader_identity_terms),
        endpoint_terms=frozenset(endpoint_terms),
        codebook_path=display_output_path(codebook_path),
        codebook_status=codebook_status,
        codebook_rows_loaded=codebook_rows_loaded,
        call_rows_loaded=len(loader_rows),
    )


def load_binding_loader_rules(codebook_path: Path | None) -> list[BindingLoaderRule]:
    if not codebook_path or not codebook_path.exists():
        return []
    try:
        from openpyxl import load_workbook
    except Exception:
        return []
    rules: list[BindingLoaderRule] = []
    try:
        wb = load_workbook(codebook_path, read_only=True, data_only=True)
        for sheet in wb.worksheets:
            if sheet.title != "model_loading":
                continue
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            header_map = {
                normalize_header(value): index
                for index, value in enumerate(header)
                if value is not None
            }
            required = {"import origin", "model loader", "terminal_call", "chain_suffix", "model argument(s)"}
            if not required.issubset(header_map):
                continue
            for values in rows:
                row = {
                    name: clean_xlsx_value(values[index] if index < len(values) else "")
                    for name, index in header_map.items()
                }
                if not row.get("import origin") or not row.get("terminal_call"):
                    continue
                rules.append(
                    BindingLoaderRule(
                        import_origin=row.get("import origin", ""),
                        model_loader=row.get("model loader", ""),
                        terminal_call=row.get("terminal_call", ""),
                        chain_suffix=row.get("chain_suffix", ""),
                        model_args=frozenset(parse_model_arg_list(row.get("model argument(s)", ""))),
                    )
                )
    except Exception:
        return []
    return rules


def parse_model_arg_list(value: str) -> set[str]:
    args = set(MODEL_ARG_NAMES)
    for part in re.split(r"[|,;/\s]+", value or ""):
        cleaned = part.strip().strip("`'\"")
        if cleaned:
            args.add(cleaned)
    return args


def jedi_sys_paths(repo: Path) -> list[str]:
    paths = [str(repo)]
    for child in ("src", "lib", "python"):
        path = repo / child
        if path.exists():
            paths.append(str(path))
    for path in repo.rglob("site-packages"):
        if path.is_dir():
            paths.append(str(path))
    return paths


def callable_node_key(facts: FileFacts, node: ast.AST) -> tuple[str, int, int]:
    return (facts.rel_path, getattr(node, "lineno", 0), getattr(node, "col_offset", 0))


def jedi_lookup_column(facts: FileFacts, node: ast.AST) -> int:
    if isinstance(node, ast.Name):
        return node.col_offset + 1
    if isinstance(node, ast.Attribute):
        line = line_at(facts.lines, node.lineno)
        segment = line[node.col_offset:]
        offset = segment.rfind(node.attr)
        if offset >= 0:
            return node.col_offset + offset + 1
    return getattr(node, "col_offset", -1) + 1


def external_import_origins(facts: FileFacts, import_origins: set[str]) -> dict[str, str]:
    origins: dict[str, str] = {}
    if facts.tree is None:
        return origins
    for node in facts.nodes:
        if isinstance(node, ast.Import):
            for alias in node.names:
                root = alias.name.split(".", 1)[0]
                if root in import_origins:
                    origins[alias.asname or root] = root
        elif isinstance(node, ast.ImportFrom) and node.module:
            root = node.module.split(".", 1)[0]
            if root not in import_origins:
                continue
            for alias in node.names:
                origins[alias.asname or alias.name] = root
    return origins


def binding_rules_for_chain(
    visible_chain: str,
    suffix_rules: dict[str, list[BindingLoaderRule]],
) -> list[BindingLoaderRule]:
    rules: list[BindingLoaderRule] = []
    for suffix, suffix_rule_rows in suffix_rules.items():
        if visible_chain == suffix or visible_chain.endswith("." + suffix):
            rules.extend(suffix_rule_rows)
    return unique_binding_rules(rules)


def unique_binding_rules(rules: Iterable[BindingLoaderRule]) -> list[BindingLoaderRule]:
    seen: set[tuple[str, str, str, str]] = set()
    output: list[BindingLoaderRule] = []
    for rule in rules:
        key = (rule.import_origin, rule.model_loader, rule.terminal_call, rule.chain_suffix)
        if key in seen:
            continue
        seen.add(key)
        output.append(rule)
    return output


def first_matching_provider_origin(
    facts: FileFacts,
    call: ast.Call,
    visible_chain: str,
    matched_rules: list[BindingLoaderRule],
    imports: dict[str, str],
) -> str:
    rule_origins = {rule.import_origin for rule in matched_rules if rule.import_origin}
    parts = visible_chain.split(".")
    if parts and imports.get(parts[0]) in rule_origins:
        return imports[parts[0]]
    if len(parts) >= 2 and parts[0] == "self":
        class_name = enclosing_class_name(facts, call)
        field_origin = class_field_provider_origin(facts, class_name, parts[1], imports, rule_origins)
        if field_origin:
            return field_origin
    return ""


def class_field_provider_origin(
    facts: FileFacts,
    class_name: str,
    field_name: str,
    imports: dict[str, str],
    allowed_origins: set[str],
) -> str:
    if not class_name:
        return ""
    for node in facts.nodes:
        if not isinstance(node, (ast.Assign, ast.AnnAssign)):
            continue
        if enclosing_class_name(facts, node) != class_name:
            continue
        targets = node.targets if isinstance(node, ast.Assign) else [node.target]
        if not any(is_self_field_target(target, field_name) for target in targets):
            continue
        value = node.value
        if not isinstance(value, ast.Call):
            continue
        chain = call_chain(value.func)
        alias = chain.split(".", 1)[0] if chain else ""
        origin = imports.get(alias, "")
        if origin in allowed_origins:
            return origin
    return ""


def enclosing_class_name(facts: FileFacts, node: ast.AST) -> str:
    cursor: ast.AST | None = node
    while cursor in facts.parents:
        cursor = facts.parents[cursor]
        if isinstance(cursor, ast.ClassDef):
            return cursor.name
    return ""


def is_self_field_target(target: ast.AST, field_name: str) -> bool:
    return (
        isinstance(target, ast.Attribute)
        and isinstance(target.value, ast.Name)
        and target.value.id in {"self", "cls"}
        and target.attr == field_name
    )


def receiver_chain_from_visible_chain(visible_chain: str) -> str:
    parts = visible_chain.split(".")
    return ".".join(parts[:-1]) if len(parts) > 1 else ""


def line_at(lines: list[str], line_number: int) -> str:
    if 1 <= line_number <= len(lines):
        return lines[line_number - 1].strip()
    return ""


def unique(values: Iterable[str]) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for value in values:
        if not value or value in seen:
            continue
        seen.add(value)
        output.append(value)
    return output


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def clean_xlsx_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def add_provider_terms(output: set[str], value: str) -> None:
    output.update(identity_terms_from_text(value, add_compact=True))


def add_loader_action_terms(output: set[str], value: str) -> None:
    output.update(identity_terms_from_text(value, add_compact=True))


def add_loader_identity_terms(output: set[str], value: str) -> None:
    output.update(identity_terms_from_text(value, add_compact=True))


def identity_terms_from_text(value: str, *, add_compact: bool = False) -> set[str]:
    text = str(value or "").strip()
    if not text:
        return set()
    terms = split_identifier(text)
    if add_compact and len(text) <= 80:
        compact = re.sub(r"[^A-Za-z0-9]+", "", text).lower()
        if compact:
            terms.add(compact)
    return terms


def terminal_from_chain(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    return text.split(".")[-1]


def sanitize_derived_terms(terms: set[str], *, keep_stopwords: bool) -> set[str]:
    output = set()
    for term in terms:
        normalized = term.strip().lower()
        if len(normalized) <= 1:
            continue
        if not keep_stopwords and normalized in DERIVED_TERM_STOPWORDS:
            continue
        output.add(normalized)
    return output


def find_literal_node_id(facts: FileFacts, literal: str, line: int, col: int) -> str:
    if facts.tree is None:
        return ""
    candidates: list[tuple[int, ast.Constant]] = []
    for node in facts.nodes:
        if isinstance(node, ast.Constant) and isinstance(node.value, str):
            if node.value != literal:
                continue
            if node.lineno != line:
                continue
            distance = abs(getattr(node, "col_offset", 0) - col)
            candidates.append((distance, node))
    if not candidates:
        return ""
    _, node = sorted(candidates, key=lambda item: item[0])[0]
    return literal_node(facts.rel_path, node.lineno, node.col_offset, literal)


def find_call_at_line(facts: FileFacts, line: int, visible_chain: str) -> ast.Call | None:
    if facts.tree is None:
        return None
    calls = [node for node in facts.nodes if isinstance(node, ast.Call) and node.lineno == line]
    if not calls:
        return None
    if visible_chain:
        for call in calls:
            if call_chain(call.func).endswith(visible_chain) or visible_chain.endswith(call_chain(call.func)):
                return call
    return calls[0]


def call_chain(node: ast.AST) -> str:
    if isinstance(node, ast.Name):
        return node.id
    if isinstance(node, ast.Attribute):
        prefix = call_chain(node.value)
        return f"{prefix}.{node.attr}" if prefix else node.attr
    if isinstance(node, ast.Call):
        return call_chain(node.func)
    return ""


def build_mock_patch_index(repo_facts: RepoFacts) -> MockPatchIndex:
    fixture_targets: dict[str, set[str]] = defaultdict(set)
    autouse_targets: set[str] = set()
    for facts in repo_facts.files.values():
        if facts.tree is None:
            continue
        for node in facts.nodes:
            if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                continue
            fixture_name, autouse = pytest_fixture_info(node)
            if not fixture_name:
                continue
            targets = patched_targets_for_root(node)
            if not targets:
                continue
            fixture_targets[fixture_name].update(targets)
            if autouse:
                autouse_targets.update(targets)
    return MockPatchIndex(
        fixture_targets={name: frozenset(targets) for name, targets in fixture_targets.items()},
        autouse_targets=frozenset(autouse_targets),
    )


def pytest_fixture_info(node: ast.FunctionDef | ast.AsyncFunctionDef) -> tuple[str, bool]:
    for decorator in node.decorator_list:
        target = decorator
        if isinstance(decorator, ast.Call):
            target = decorator.func
        chain = call_chain(target)
        if chain not in {"fixture", "pytest.fixture"} and not chain.endswith(".fixture"):
            continue
        fixture_name = node.name
        autouse = False
        if isinstance(decorator, ast.Call):
            for keyword in decorator.keywords:
                if keyword.arg == "name":
                    explicit_name = literal_string(keyword.value)
                    if explicit_name:
                        fixture_name = explicit_name
                elif keyword.arg == "autouse" and isinstance(keyword.value, ast.Constant):
                    autouse = bool(keyword.value.value)
        return fixture_name, autouse
    return "", False


def mocked_loader_reason(
    facts: FileFacts,
    call: ast.Call,
    sink_row: dict[str, str],
    identity_vocab: IdentityVocabulary,
    mock_patch_index: MockPatchIndex,
) -> str:
    chain = call_chain(call.func).lower()
    signals = []
    if any(term in chain for term in MOCK_SYMBOL_TERMS):
        signals.append("mocked_loader_symbol")
    patched_targets = patched_targets_for_node(facts, call, mock_patch_index)
    matched_targets = [
        target for target in patched_targets
        if patch_target_matches_sink(target, chain, sink_row, identity_vocab)
    ]
    if matched_targets:
        signals.append("mock_patch_target:" + ",".join(sorted(matched_targets)[:4]))
    return ",".join(sorted(set(signals)))


def patched_targets_for_node(
    facts: FileFacts,
    node: ast.AST,
    mock_patch_index: MockPatchIndex | None = None,
) -> set[str]:
    root = enclosing_function(node, facts.parents) or facts.tree
    targets = patched_targets_for_root(root)
    class_root = enclosing_class(node, facts.parents)
    targets.update(class_context_patch_targets(class_root))
    if mock_patch_index:
        targets.update(fixture_patch_targets_for_node(facts, node, mock_patch_index))
    return targets


def source_mocked_reason(builder: SourceSinkBuilder, source_row: dict[str, str], sink: SinkInfo) -> str:
    rel_path = source_row.get("file_path", "")
    line = parse_int(source_row.get("line_number"))
    facts = builder.repo_facts.files.get(rel_path)
    if not facts or line <= 0:
        return ""
    root = enclosing_function_for_line(facts, line) or facts.tree
    patched_targets = patched_targets_for_root(root)
    class_root = enclosing_class_for_line(facts, line)
    patched_targets.update(class_context_patch_targets(class_root))
    if root is not None:
        patched_targets.update(fixture_patch_targets_for_root(facts, root, builder.mock_patch_index))
    if is_test_path(facts.rel_path):
        patched_targets.update(builder.mock_patch_index.autouse_targets)
    if not patched_targets:
        return ""
    chain = sink.loader_row.get("visible_call_chain", "")
    matched_targets = [
        target for target in patched_targets
        if patch_target_matches_sink(target, chain, sink.loader_row, builder.identity_vocab)
    ]
    signals = []
    if matched_targets:
        signals.append("source_scope_mock_patch_target:" + ",".join(sorted(matched_targets)[:4]))
    return ",".join(sorted(set(signals)))


def patched_targets_for_root(root: ast.AST | None) -> set[str]:
    if root is None:
        return set()
    targets: set[str] = set()
    for candidate in ast.walk(root):
        if isinstance(candidate, ast.Call):
            targets.update(mock_patch_targets_from_call(candidate))
        elif isinstance(candidate, (ast.Assign, ast.AnnAssign)):
            value = candidate.value
            if not value_looks_mock_replacement(value):
                continue
            assign_targets = candidate.targets if isinstance(candidate, ast.Assign) else [candidate.target]
            for target in assign_targets:
                target_chain = call_chain(target)
                if target_chain:
                    targets.add(target_chain)
    return targets


def class_context_patch_targets(root: ast.ClassDef | None) -> set[str]:
    if root is None:
        return set()
    targets: set[str] = set()
    for decorator in root.decorator_list:
        if isinstance(decorator, ast.Call):
            targets.update(mock_patch_targets_from_call(decorator))
    setup_names = {"setUp", "setUpClass", "asyncSetUp", "asyncSetUpClass"}
    for item in root.body:
        if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)) and item.name in setup_names:
            targets.update(patched_targets_for_root(item))
    return targets


def fixture_patch_targets_for_node(
    facts: FileFacts,
    node: ast.AST,
    mock_patch_index: MockPatchIndex,
) -> set[str]:
    root = enclosing_function(node, facts.parents)
    return fixture_patch_targets_for_root(facts, root, mock_patch_index)


def fixture_patch_targets_for_root(
    facts: FileFacts,
    root: ast.AST | None,
    mock_patch_index: MockPatchIndex,
) -> set[str]:
    if not isinstance(root, (ast.FunctionDef, ast.AsyncFunctionDef)):
        return set()
    targets: set[str] = set()
    arg_names = function_arg_names(root)
    for arg_name in arg_names:
        targets.update(mock_patch_index.fixture_targets.get(arg_name, frozenset()))
    if is_test_path(facts.rel_path):
        targets.update(mock_patch_index.autouse_targets)
    return targets


def enclosing_function_for_line(facts: FileFacts, line: int) -> ast.AST | None:
    if facts.tree is None:
        return None
    candidates = []
    for node in facts.nodes:
        if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            continue
        start = getattr(node, "lineno", 0)
        end = getattr(node, "end_lineno", start)
        if start <= line <= end:
            candidates.append((end - start, node))
    if not candidates:
        return None
    return sorted(candidates, key=lambda item: item[0])[0][1]


def enclosing_class_for_line(facts: FileFacts, line: int) -> ast.ClassDef | None:
    if facts.tree is None:
        return None
    candidates = []
    for node in facts.nodes:
        if not isinstance(node, ast.ClassDef):
            continue
        start = getattr(node, "lineno", 0)
        end = getattr(node, "end_lineno", start)
        if start <= line <= end:
            candidates.append((end - start, node))
    if not candidates:
        return None
    return sorted(candidates, key=lambda item: item[0])[0][1]


def mock_patch_targets_from_call(call: ast.Call) -> set[str]:
    chain = call_chain(call.func)
    normalized_chain = chain.lower()
    targets: set[str] = set()
    if normalized_chain.endswith("patch") and call.args:
        target = literal_string(call.args[0])
        if target:
            targets.add(target)
    if normalized_chain.endswith("patch.object") and len(call.args) >= 2:
        attr = literal_string(call.args[1])
        if attr:
            targets.add(f"{expr_text(call.args[0])}.{attr}")
    if "monkeypatch" in normalized_chain and normalized_chain.endswith("setattr") and call.args:
        if len(call.args) >= 2:
            dotted_target = literal_string(call.args[0])
            attr = literal_string(call.args[1])
            if dotted_target and not attr:
                targets.add(dotted_target)
            elif attr:
                targets.add(f"{expr_text(call.args[0])}.{attr}")
    return {target for target in targets if target}


def value_looks_mock_replacement(value: ast.AST | None) -> bool:
    if value is None:
        return False
    if isinstance(value, ast.Lambda):
        return True
    text = expr_text(value).lower()
    return any(term in text for term in MOCK_SYMBOL_TERMS) or "magicmock" in text or "mock(" in text


def patch_target_matches_call(target: str, chain: str, identity_vocab: IdentityVocabulary) -> bool:
    target_norm = normalize_dotted_name(target)
    chain_norm = normalize_dotted_name(chain)
    if not target_norm or not chain_norm:
        return False
    if target_norm.endswith(chain_norm) or chain_norm.endswith(target_norm):
        return True
    if chain_norm.startswith(target_norm + "."):
        return True
    target_parts = dotted_parts(target_norm)
    chain_parts = dotted_parts(chain_norm)
    if not target_parts or not chain_parts:
        return False
    if chain_parts[-1] in target_parts and (set(target_parts) & {"init", "new"}):
        return True
    if target_parts[-1] != chain_parts[-1]:
        return False
    target_set = set(target_parts)
    chain_set = set(chain_parts)
    provider_hit = bool(target_set & set(identity_vocab.provider_terms))
    loader_overlap = (target_set & chain_set) & set(identity_vocab.endpoint_terms)
    return provider_hit and len(loader_overlap) >= 1


def patch_target_matches_sink(
    target: str,
    chain: str,
    sink_row: dict[str, str],
    identity_vocab: IdentityVocabulary,
) -> bool:
    target_norm = normalize_dotted_name(target)
    if not target_norm:
        return False
    chain_norm = normalize_dotted_name(chain)
    identity_values = [
        chain,
        sink_row.get("visible_call_chain", ""),
        sink_row.get("terminal_call", ""),
        sink_row.get("receiver_symbol", ""),
    ]
    identity_values.extend((sink_row.get("matched_chain_suffix", "") or "").split("|"))
    identity_values.extend((sink_row.get("matched_canonical_loader", "") or "").split("|"))
    for value in identity_values:
        value_norm = normalize_dotted_name(value)
        if not value_norm:
            continue
        if target_norm.endswith(value_norm) or value_norm.endswith(target_norm):
            return True
        if value_norm.startswith(target_norm + "."):
            return True

    target_parts = set(dotted_parts(target_norm))
    if not target_parts:
        return False
    sink_origin = (
        sink_row.get("receiver_origin")
        or sink_row.get("linked_import_origin")
        or sink_row.get("matched_rule_import_origin")
    )
    origin_terms = split_identifier(sink_origin.replace(".", "_")) if sink_origin else set()
    if origin_terms and not (target_parts & origin_terms):
        return False
    sink_terms: set[str] = set()
    for value in identity_values:
        sink_terms.update(dotted_parts(normalize_dotted_name(value)))
    endpoint_overlap = (target_parts & sink_terms) & set(identity_vocab.endpoint_terms)
    loader_overlap = target_parts & sink_terms & set(identity_vocab.loader_identity_terms)
    return bool(origin_terms) and bool(endpoint_overlap or loader_overlap)


def is_provider_endpoint_patch(target: str, identity_vocab: IdentityVocabulary) -> bool:
    parts = set(dotted_parts(normalize_dotted_name(target)))
    return bool(parts & set(identity_vocab.provider_terms)) and bool(parts & set(identity_vocab.endpoint_terms))


def is_test_path(rel_path: str) -> bool:
    parts = {part.lower() for part in re.split(r"[/_.-]+", rel_path)}
    return bool(parts & {"test", "tests", "testing"})


def normalize_dotted_name(value: str) -> str:
    return ".".join(part for part in re.split(r"[^A-Za-z0-9_]+", value.lower()) if part)


def dotted_parts(value: str) -> list[str]:
    parts = []
    for part in value.split("."):
        split = split_identifier(part)
        parts.extend(sorted(split) if split else [part])
    return [part for part in parts if part]


def enclosing_function(node: ast.AST, parents: dict[ast.AST, ast.AST]) -> ast.AST | None:
    cursor = node
    while cursor in parents:
        cursor = parents[cursor]
        if isinstance(cursor, (ast.FunctionDef, ast.AsyncFunctionDef)):
            return cursor
    return None


def enclosing_class(node: ast.AST, parents: dict[ast.AST, ast.AST]) -> ast.ClassDef | None:
    cursor = node
    while cursor in parents:
        cursor = parents[cursor]
        if isinstance(cursor, ast.ClassDef):
            return cursor
    return None


def function_arg_names(node: ast.FunctionDef | ast.AsyncFunctionDef) -> set[str]:
    args = node.args
    names = {
        arg.arg
        for arg in args.posonlyargs + args.args + args.kwonlyargs
        if arg.arg not in {"self", "cls"}
    }
    if args.vararg and args.vararg.arg not in {"self", "cls"}:
        names.add(args.vararg.arg)
    if args.kwarg and args.kwarg.arg not in {"self", "cls"}:
        names.add(args.kwarg.arg)
    return names


def is_model_arg_name(name: str | None) -> bool:
    if not name:
        return False
    normalized = name.strip().lower()
    if normalized in MODEL_ARG_NAMES or normalized in MODEL_PAYLOAD_KEYS:
        return True
    terms = split_identifier(normalized)
    return "model" in terms and not terms & {"metadata", "registry", "list", "types"}


def class_looks_like_model_registry(info: ClassDefInfo) -> bool:
    terms = split_identifier(info.name)
    base_terms = {term for base in info.bases for term in split_identifier(base)}
    if "enum" not in base_terms and "model" not in terms:
        return False
    for node in info.node.body:
        value = None
        if isinstance(node, ast.Assign):
            value = node.value
        elif isinstance(node, ast.AnnAssign):
            value = node.value
        if isinstance(value, ast.Constant) and isinstance(value.value, str):
            return True
    return False


def class_looks_like_config(info: ClassDefInfo) -> bool:
    terms = split_identifier(info.name)
    base_terms = {term for base in info.bases for term in split_identifier(base)}
    return bool((terms | base_terms) & GENERIC_CONFIG_TERMS)


def expr_reads_config_model_field(expr: ast.AST | None) -> bool:
    if expr is None:
        return False
    for node in ast.walk(expr):
        if isinstance(node, ast.Attribute) and is_model_arg_name(node.attr):
            receiver_terms = split_identifier(expr_text(node.value))
            if receiver_terms & GENERIC_CONFIG_TERMS:
                return True
        if isinstance(node, ast.Subscript) and subscript_key(node) in MODEL_PAYLOAD_KEYS:
            receiver_terms = split_identifier(expr_text(node.value))
            if receiver_terms & GENERIC_CONFIG_TERMS:
                return True
    return False


def is_config_arg_name(name: str | None) -> bool:
    if not name:
        return False
    normalized = name.strip().lower()
    if normalized in GENERIC_CONFIG_ARG_NAMES:
        return True
    terms = split_identifier(normalized)
    return bool(terms & GENERIC_CONFIG_TERMS)


def expr_looks_like_config_arg(expr: ast.AST) -> bool:
    if isinstance(expr, ast.Dict):
        return True
    return bool(split_identifier(expr_text(expr)) & GENERIC_CONFIG_TERMS)


def terms_from_labels(labels: Iterable[str]) -> set[str]:
    terms: set[str] = set()
    for label in labels:
        terms.update(split_identifier(label))
    return terms


def split_identifier(value: str) -> set[str]:
    value = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1_\2", value)
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", value)
    return {
        part.lower()
        for part in re.split(r"[^A-Za-z0-9]+", value)
        if part
    }


def target_names(target: ast.AST) -> list[str]:
    if isinstance(target, ast.Name):
        return [target.id]
    if isinstance(target, (ast.Tuple, ast.List)):
        output = []
        for item in target.elts:
            output.extend(target_names(item))
        return output
    return []


def target_text(target: ast.AST) -> str:
    return expr_text(target)


def expr_text(node: ast.AST | None) -> str:
    if node is None:
        return ""
    try:
        return ast.unparse(node)
    except Exception:
        return node.__class__.__name__


def literal_key(node: ast.AST | None) -> str:
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        return node.value
    return ""


def literal_string(node: ast.AST | None) -> str:
    return literal_key(node)


def subscript_key(node: ast.Subscript) -> str:
    key = node.slice
    if isinstance(key, ast.Constant) and isinstance(key.value, str):
        return key.value
    return ""


def infer_repo_snapshot(output_dir: Path) -> str:
    return output_dir.name


def location_url_from_loader(loader_row: dict[str, str], source_row: dict[str, str]) -> str:
    loader_url = loader_row.get("location_url", "")
    if not loader_url:
        return ""
    prefix = loader_url.split("/blob/")[0]
    after_blob = loader_url.split("/blob/", 1)[1]
    commit = after_blob.split("/", 1)[0]
    return f"{prefix}/blob/{commit}/{source_row.get('file_path', '')}#L{source_row.get('line_number', '')}"


def node_file_part(node: str) -> str:
    parts = node.split("|")
    for part in parts:
        if part.endswith(".py") or "/" in part:
            return part
    return ""


def var_node(scope_id: str, name: str) -> str:
    return f"var|{scope_id}|{name}"


def field_node(class_symbol: str, attr: str) -> str:
    return f"field|{class_symbol}|{attr}"


def payload_node(scope_id: str, name: str) -> str:
    return f"payload|{scope_id}|{name}|model"


def payload_version_node(scope_id: str, name: str, line_number: int) -> str:
    return f"payload_version|{scope_id}|{name}|{line_number}|model"


def payload_literal_node(rel_path: str, line: int, col: int, key: str) -> str:
    return f"payload_literal|{rel_path}|{line}|{col}|{key}"


def mapping_item_node(mapping_node: str, key: str) -> str:
    return f"mapping_item|{mapping_node}|{key}"


def payload_key_from_node(node: str) -> str:
    parts = node.split("|")
    if not parts:
        return ""
    if parts[0] == "payload_literal" and len(parts) >= 5:
        return parts[-1]
    if parts[0] == "payload" and len(parts) >= 4:
        return parts[-2]
    return ""


def literal_node(rel_path: str, line: int, col: int, value: str) -> str:
    return f"literal|{rel_path}|{line}|{col}|{value}"


def config_value_node(rel_path: str, line: int, key: str, value: str) -> str:
    return f"config_value|{rel_path}|{line}|{key}|{value}"


def call_result_node(rel_path: str, line: int, col: int) -> str:
    return f"call_result|{rel_path}|{line}|{col}"


def return_node(function_symbol: str) -> str:
    return f"return|{function_symbol}"


def sink_node(loader_candidate_id: str, rel_path: str, line: int) -> str:
    return f"sink|{loader_candidate_id}|{rel_path}|{line}"


def module_node(module: str) -> str:
    return f"module|{module}"


def attr_node(base: str, attr: str) -> str:
    return f"attr|{base}|{attr}"


def widget_node(class_symbol: str, widget_attr: str) -> str:
    return f"widget|{class_symbol}|{widget_attr}|currentText"


def config_key_from_line(line_text: str) -> str:
    match = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*[:=]", line_text or "")
    if not match:
        return ""
    key = match.group(1)
    terms = split_identifier(key)
    if is_model_arg_name(key) or terms & {"llm", "model", "deployment", "engine", "embedder", "embedding", "reranker"}:
        return key
    return ""
