"""Load terms for candidate calls and PTM ID source context."""

from __future__ import annotations

from mist.resources import DATA_DIR

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any


DEFAULT_CONTEXT_TERMS = DATA_DIR / "context_terms.json"
DEFAULT_REUSE_CODEBOOK = DATA_DIR / "reuse_rules.xlsx"

FALLBACK_CALL_GENERIC_TERMINALS = frozenset({
    "call",
    "complete",
    "create",
    "generate",
    "open",
    "post",
    "request",
    "send",
})

FALLBACK_CALL_MODEL_ARGS = frozenset({
    "model",
    "model_name",
    "deployment_name",
    "embedding_model",
    "small_model",
})

FALLBACK_CONTEXT_MODEL_CARRIER_TERMS = frozenset({
    "model",
    "models",
    "model_id",
    "model_name",
    "model_type",
    "deployment",
    "deployment_id",
    "deployment_name",
    "engine",
    "llm",
    "chat_model",
    "embedding_model",
})

FALLBACK_CONTEXT_METADATA_TERMS = frozenset({
    "token",
    "tokens",
    "tokenizer",
    "counter",
    "count",
    "price",
    "pricing",
    "cost",
    "schema",
    "context",
    "context_window",
    "limit",
    "max",
    "encoding",
})

FALLBACK_CONTEXT_DESCRIPTIVE_TERMS = frozenset({
    "description",
    "desc",
    "doc",
    "docs",
    "example",
    "examples",
    "sample",
    "samples",
    "tutorial",
    "readme",
    "comment",
    "message",
})

FALLBACK_CONTEXT_MODEL_ID_VALIDATION_CALLS = frozenset({
    "any",
    "all",
})

FALLBACK_CONTEXT_MODEL_ID_VALIDATION_METHODS = frozenset({
    "startswith",
    "endswith",
})

FALLBACK_CONTEXT_FUNCTION_DEFAULT_EXTRA_MODEL_TERMS = frozenset({
    "orchestrator",
})

FALLBACK_CONTEXT_CLI_COMMAND_TARGET_TERMS = frozenset({
    "args",
    "argv",
    "cmd",
    "command",
    "commands",
})

FALLBACK_CONTEXT_SHORT_AMBIGUOUS_OPENAI_MODEL_IDS = frozenset({
    "o1",
    "o3",
})


@dataclass(frozen=True)
class CallTermConfig:
    """call analysis loader-candidate matching terms."""

    model_args: frozenset[str]
    generic_terminals: frozenset[str]
    codebook_model_args: frozenset[str]


@dataclass(frozen=True)
class ContextTermConfig:
    """source context analysis local-context routing terms."""

    model_carrier_terms: frozenset[str]
    metadata_terms: frozenset[str]
    descriptive_terms: frozenset[str]
    model_id_validation_calls: frozenset[str]
    model_id_validation_methods: frozenset[str]
    function_default_model_terms: frozenset[str]
    cli_command_target_terms: frozenset[str]
    short_ambiguous_openai_model_ids: frozenset[str]


def build_call_term_config(
    reuse_codebook: Path | None = DEFAULT_REUSE_CODEBOOK,
    context_terms_path: Path | None = DEFAULT_CONTEXT_TERMS,
) -> CallTermConfig:
    """Build call analysis terms from the context config and reuse codebook."""

    config = load_context_terms(context_terms_path)
    call = config.get("call", {})
    fallback_model_args = config_terms(call, "fallback_model_args", FALLBACK_CALL_MODEL_ARGS)
    codebook_model_args = frozenset(read_reuse_codebook_model_args(reuse_codebook))
    return CallTermConfig(
        model_args=frozenset(fallback_model_args | codebook_model_args),
        generic_terminals=config_terms(call, "generic_terminals", FALLBACK_CALL_GENERIC_TERMINALS),
        codebook_model_args=codebook_model_args,
    )


def build_context_term_config(
    context_terms_path: Path | None = DEFAULT_CONTEXT_TERMS,
) -> ContextTermConfig:
    """Build source context analysis context terms from the external codebook config."""

    config = load_context_terms(context_terms_path)
    context = config.get("context", {})
    model_carrier_terms = config_terms(
        context,
        "model_carrier_terms",
        FALLBACK_CONTEXT_MODEL_CARRIER_TERMS,
    )
    function_default_extra_terms = config_terms(
        context,
        "function_default_extra_model_terms",
        FALLBACK_CONTEXT_FUNCTION_DEFAULT_EXTRA_MODEL_TERMS,
    )
    return ContextTermConfig(
        model_carrier_terms=model_carrier_terms,
        metadata_terms=config_terms(context, "metadata_terms", FALLBACK_CONTEXT_METADATA_TERMS),
        descriptive_terms=config_terms(context, "descriptive_terms", FALLBACK_CONTEXT_DESCRIPTIVE_TERMS),
        model_id_validation_calls=config_terms(
            context,
            "model_id_validation_calls",
            FALLBACK_CONTEXT_MODEL_ID_VALIDATION_CALLS,
        ),
        model_id_validation_methods=config_terms(
            context,
            "model_id_validation_methods",
            FALLBACK_CONTEXT_MODEL_ID_VALIDATION_METHODS,
        ),
        function_default_model_terms=frozenset(model_carrier_terms | function_default_extra_terms),
        cli_command_target_terms=config_terms(
            context,
            "cli_command_target_terms",
            FALLBACK_CONTEXT_CLI_COMMAND_TARGET_TERMS,
        ),
        short_ambiguous_openai_model_ids=config_terms(
            context,
            "short_ambiguous_openai_model_ids",
            FALLBACK_CONTEXT_SHORT_AMBIGUOUS_OPENAI_MODEL_IDS,
        ),
    )


def read_reuse_codebook_model_args(reuse_codebook: Path | None) -> set[str]:
    """Derive plain model-argument names from `model argument(s)` cells."""

    if not reuse_codebook or not reuse_codebook.exists():
        return set()
    try:
        from openpyxl import load_workbook
    except Exception:
        return set()

    output: set[str] = set()
    try:
        workbook = load_workbook(reuse_codebook, read_only=True, data_only=True)
        worksheet = workbook["model_loading"]
        headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
        if "model argument(s)" not in headers:
            return set()
        column = headers.index("model argument(s)") + 1
        for row in range(2, worksheet.max_row + 1):
            output.update(parse_model_arg_tokens(worksheet.cell(row, column).value))
    except Exception:
        return set()
    return output


def parse_model_arg_tokens(value: Any) -> set[str]:
    """Normalize codebook model-argument cells into plain Python-ish names."""

    output: set[str] = set()
    text = str(value or "")
    for piece in text.replace("/", " ").replace(",", " ").replace(";", " ").split():
        cleaned = piece.strip()
        if not cleaned:
            continue
        if "." in cleaned:
            cleaned = cleaned.rsplit(".", 1)[-1]
        cleaned = cleaned.strip("[]`'\"")
        if cleaned:
            output.add(cleaned)
    return output


def load_context_terms(path: Path | None) -> dict[str, Any]:
    """Load context terms, using defaults if the config is absent."""

    if not path or not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def config_terms(config: dict[str, Any], key: str, fallback: frozenset[str]) -> frozenset[str]:
    """Return normalized terms from config, or a stable fallback set."""

    raw = config.get(key)
    if not isinstance(raw, list):
        return fallback
    output = {
        normalize_term(value)
        for value in raw
        if normalize_term(value)
    }
    return frozenset(output) if output else fallback


def normalize_term(value: Any) -> str:
    return re.sub(r"\s+", "_", str(value or "").strip().lower())


CALL_TERM_CONFIG = build_call_term_config()
CONTEXT_TERM_CONFIG = build_context_term_config()

CALL_GENERIC_TERMINALS = CALL_TERM_CONFIG.generic_terminals
CALL_DEFAULT_MODEL_ARGS = CALL_TERM_CONFIG.model_args

CONTEXT_MODEL_CARRIER_TERMS = CONTEXT_TERM_CONFIG.model_carrier_terms
CONTEXT_METADATA_TERMS = CONTEXT_TERM_CONFIG.metadata_terms
CONTEXT_DESCRIPTIVE_TERMS = CONTEXT_TERM_CONFIG.descriptive_terms
CONTEXT_MODEL_ID_VALIDATION_CALLS = CONTEXT_TERM_CONFIG.model_id_validation_calls
CONTEXT_MODEL_ID_VALIDATION_METHODS = CONTEXT_TERM_CONFIG.model_id_validation_methods
CONTEXT_FUNCTION_DEFAULT_MODEL_TERMS = CONTEXT_TERM_CONFIG.function_default_model_terms
CONTEXT_CLI_COMMAND_TARGET_TERMS = CONTEXT_TERM_CONFIG.cli_command_target_terms
CONTEXT_SHORT_AMBIGUOUS_OPENAI_MODEL_IDS = CONTEXT_TERM_CONFIG.short_ambiguous_openai_model_ids
