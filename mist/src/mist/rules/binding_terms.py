"""Terms used to connect PTM IDs to eligible calls."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any


DEFAULT_MODEL_ARG_NAMES = frozenset({
    "model",
    "model_name",
    "model_id",
    "deployment",
    "deployment_id",
    "deployment_name",
    "engine",
    "llm",
    "chat_model",
    "embedding_model",
    "small_model",
})

DEFAULT_MODEL_PAYLOAD_KEYS = frozenset({
    "model",
    "model_name",
    "model_id",
    "deployment",
    "deployment_name",
    "engine",
    "embedding_model",
    "small_model",
})

REQUEST_BODY_ARG_NAMES = frozenset({
    "body",
    "data",
    "json",
})

VALUE_PRESERVING_METHODS = frozenset({
    "casefold",
    "copy",
    "decode",
    "encode",
    "lower",
    "strip",
    "upper",
})

GENERIC_CONFIG_ARG_NAMES = frozenset({
    "config",
    "configs",
    "configuration",
    "settings",
    "options",
    "params",
    "parameters",
    "kwargs",
})

GENERIC_CONFIG_TERMS = GENERIC_CONFIG_ARG_NAMES | frozenset({"kwarg", "argument", "arguments"})

GENERIC_ROLE_TERMS = frozenset({
    "adapter",
    "adapters",
    "client",
    "clients",
    "factory",
    "loader",
    "manager",
    "proxy",
    "registry",
    "router",
    "service",
    "wrapper",
})

GENERIC_ENDPOINT_TERMS = frozenset({
    "call",
    "chat",
    "completion",
    "completions",
    "create",
    "generate",
    "invoke",
    "post",
    "request",
    "response",
    "responses",
    "run",
})

DERIVED_TERM_STOPWORDS = frozenset({
    "ai",
    "api",
    "aio",
    "async",
    "base",
    "batch",
    "class",
    "client",
    "clients",
    "cls",
    "component",
    "content",
    "core",
    "ext",
    "from",
    "function",
    "get",
    "index",
    "init",
    "interface",
    "module",
    "model",
    "models",
    "object",
    "provider",
    "providers",
    "py",
    "python",
    "sdk",
    "set",
    "util",
    "utils",
})

MOCK_SYMBOL_TERMS = frozenset({
    "mock",
    "fake",
    "stub",
})


@dataclass(frozen=True)
class BindingTermConfig:
    """Terms used by binding tracing plus raw model args derived from the codebook."""

    model_arg_names: frozenset[str] = DEFAULT_MODEL_ARG_NAMES
    model_payload_keys: frozenset[str] = DEFAULT_MODEL_PAYLOAD_KEYS
    codebook_model_arg_names: frozenset[str] = frozenset()


def build_binding_term_config(codebook_path: Path | None) -> BindingTermConfig:
    """Load keyword terms and retain codebook argument paths separately.

    Nested paths such as ``requests[].params.model`` are not plain Python
    keyword names and are not added to the active keyword sets.
    """

    return BindingTermConfig(
        model_arg_names=DEFAULT_MODEL_ARG_NAMES,
        model_payload_keys=DEFAULT_MODEL_PAYLOAD_KEYS,
        codebook_model_arg_names=frozenset(read_codebook_model_arg_names(codebook_path)),
    )


def read_codebook_model_arg_names(codebook_path: Path | None) -> set[str]:
    """Read raw ``model argument(s)`` values from the reuse codebook."""

    if not codebook_path or not codebook_path.exists():
        return set()
    try:
        from openpyxl import load_workbook
    except Exception:
        return set()

    output: set[str] = set()
    try:
        wb = load_workbook(codebook_path, read_only=True, data_only=True)
        for sheet in wb.worksheets:
            if sheet.title != "model_loading":
                continue
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            header_map = {
                _normalize_header(value): index
                for index, value in enumerate(header)
                if value is not None
            }
            index = header_map.get("model argument(s)")
            if index is None:
                continue
            for values in rows:
                value = values[index] if index < len(values) else ""
                for part in re.split(r"[|,;/\s]+", str(value or "")):
                    cleaned = part.strip().strip("`'\"")
                    if cleaned:
                        output.add(cleaned)
    except Exception:
        return set()
    return output


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())
