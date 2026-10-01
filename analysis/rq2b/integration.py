"""Summarize reviewed integration changes and affected binding-path locations."""

import argparse
from collections import Counter, defaultdict
from contextlib import closing
import csv
import json
from pathlib import Path
import sqlite3
from statistics import median

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from analyze import HERE, read_database, read_counts, sha256, share


def read_csv(path):
    with path.open(encoding='utf-8', newline='') as stream:
        return list(csv.DictReader(stream))


def read_annotations(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        tables = {}
        for name in ('review_decisions', 'taxonomy'):
            rows = list(book[name].values)
            if not rows or any(cell.data_type == 'f' for row in book[name] for cell in row):
                raise ValueError('Missing annotation rows or unexpected formulas')
            tables[name] = [dict(zip(rows[0], row)) for row in rows[1:] if any(row)]
    finally:
        book.close()
    taxonomy = [{key: row[key] for key in ('ID', 'Change category', 'Description')}
                for row in tables['taxonomy']]
    code_ids = {row['ID'] for row in taxonomy}
    if not code_ids or len(code_ids) != len(taxonomy) or any(not row['Description'] for row in taxonomy):
        raise ValueError('Empty or repeated taxonomy codes')
    cases = {}
    for row in tables['review_decisions']:
        cid = row['Case']
        codes = [s.strip() for s in (row['Change codes'] or '').split(',') if s.strip()]
        if not cid or cid in cases or row['Outcome'] not in ('change', 'unchanged'):
            raise ValueError('Invalid or repeated annotation case')
        if len(codes) != len(set(codes)) or not set(codes) <= code_ids:
            raise ValueError('Unknown or repeated change code: ' + cid)
        if bool(codes) != (row['Outcome'] == 'change'):
            raise ValueError('Change codes do not match outcome: ' + cid)
        cases[cid] = dict(row, codes=set(codes))
    return cases, taxonomy


def read_matches(database, pairs, counts):
    stable = {r['pair_id'] for r in pairs if counts[r['old_release_id']] == counts[r['new_release_id']]}
    with closing(sqlite3.connect(database.resolve().as_uri() + '?mode=ro', uri=True)) as db:
        db.row_factory = sqlite3.Row
        rows = [dict(row) for row in db.execute('''
            SELECT m.match_id,m.pair_id,old.ptm_id AS old_ptm_id,new.ptm_id AS new_ptm_id
            FROM binding_matches m
            JOIN bindings ob ON ob.binding_id=m.old_binding_id
            JOIN occurrences old ON old.occurrence_id=ob.occurrence_id
            JOIN bindings nb ON nb.binding_id=m.new_binding_id
            JOIN occurrences new ON new.occurrence_id=nb.occurrence_id
        ''') if row['pair_id'] in stable]
    if any(r['old_ptm_id'] != r['new_ptm_id'] for r in rows):
        raise ValueError('Continuing bindings have different PTM IDs')
    return {row['match_id']: row for row in rows}


def check_selection(screens, annotations, matches, pairs, releases):
    by_case = {}
    if len({r['match_id'] for r in screens}) != len(screens):
        raise ValueError('Repeated match in screening input')
    if {r['match_id'] for r in screens} != set(matches):
        raise ValueError('Screening does not cover exactly the stable-count matched population')
    for row in screens:
        cid, status, outcome = row['case_id'], row['screen_status'], row['review_outcome']
        if not cid or cid in by_case or status not in ('edited', 'no_edit_found', 'unresolved'):
            raise ValueError('Invalid or repeated screening case')
        if status == 'edited':
            if outcome not in ('change', 'unchanged', 'invalid', 'unresolved'):
                raise ValueError('Edited case lacks a review outcome: ' + cid)
        elif outcome:
            raise ValueError('Unselected screening case has a review outcome: ' + cid)
        match = matches[row['match_id']]
        by_case[cid] = dict(row, pair_id=match['pair_id'])
    valid = {cid for cid, row in by_case.items() if row['review_outcome'] in ('change', 'unchanged')}
    if valid != set(annotations):
        raise ValueError('Annotation workbook does not cover exactly the valid reviews')
    for cid, row in annotations.items():
        selected = by_case[cid]
        pair = pairs[selected['pair_id']]
        old, new = releases[pair['old_release_id']], releases[pair['new_release_id']]
        if (row['Outcome'] != selected['review_outcome'] or
                row['Repository'].casefold() != old['repository'].casefold() or
                row['Release pair'] != old['tag'] + ' --> ' + new['tag'] or
                row.get('Git comparison', pair['compare_url']) != pair['compare_url']):
            raise ValueError('Annotation differs from the matched release pair: ' + cid)
    return by_case


def location_sets(rows, changed):
    sets, seen = {}, set()
    for row in rows:
        cid, entity, location = row['case_id'], row['entity'], row['location']
        key = cid, entity, location
        if cid not in changed or entity not in ('file', 'procedure') or not location or key in seen:
            raise ValueError('Unknown, empty, or repeated path location')
        flags = {field: row[field] for field in ('in_old', 'in_new', 'affected', 'moved')}
        if any(value not in ('0', '1') for value in flags.values()) or flags['in_old'] == flags['in_new'] == '0':
            raise ValueError('Invalid location flags')
        seen.add(key)
        case = sets.setdefault(cid, {kind: {s: set() for s in ('old', 'new', 'affected', 'moved')}
                                    for kind in ('file', 'procedure')})
        for flag, side in [('in_old', 'old'), ('in_new', 'new'), ('affected', 'affected'), ('moved', 'moved')]:
            if flags[flag] == '1':
                case[entity][side].add(location)
    if set(sets) != changed:
        raise ValueError('Reviewed locations do not cover exactly the changed cases')
    for cid, case in sets.items():
        for side in ('old', 'new'):
            for procedure in case['procedure'][side]:
                path, separator, name = procedure.partition('::')
                if not separator or not name or path not in case['file'][side]:
                    raise ValueError('Procedure lacks its path file: ' + cid)
    return sets


def measure(sets):
    old, new, affected, moved = (sets[key] for key in ('old', 'new', 'affected', 'moved'))
    total = old | new
    if not affected <= total or not moved <= total:
        raise ValueError('Affected or moved locations fall outside the path')
    added, removed = new - old, old - new
    surface = ('restructured' if added and removed else 'expanded' if added else
               'reduced' if removed else 'moved' if moved else 'unchanged')
    return dict(old=len(old), new=len(new), total=len(total), affected=len(affected),
                share=len(affected) / len(total) if total else None, surface=surface)


def combined_surface(values):
    if 'restructured' in values or {'expanded', 'reduced'} <= set(values):
        return 'restructured'
    return next((name for name in ('expanded', 'reduced', 'moved') if name in values), 'unchanged')


def aggregate_locations(locations, selected):
    groups = {}
    for cid, case in locations.items():
        pair = selected[cid]['pair_id']
        group = groups.setdefault(pair, {kind: {s: set() for s in ('old', 'new', 'affected', 'moved')}
                                         for kind in ('file', 'procedure')})
        for kind in group:
            for side in group[kind]:
                group[kind][side].update(case[kind][side])
    return {pair: {kind: measure(sets) for kind, sets in group.items()} for pair, group in groups.items()}


def category_counts(annotations, selected, taxonomy):
    changed = {cid: row for cid, row in annotations.items() if row['Outcome'] == 'change'}
    changed_pairs = {selected[cid]['pair_id'] for cid in changed}

    def count(codes, name, description):
        cases = {cid for cid, row in changed.items() if row['codes'] & set(codes)}
        release_pairs = {selected[cid]['pair_id'] for cid in cases}
        return dict(code=', '.join(codes), category=name, description=description,
                    binding_pairs=len(cases), binding_share=share(len(cases), len(changed)),
                    release_pairs=len(release_pairs), release_share=share(len(release_pairs), len(changed_pairs)))

    rows = [count([r['ID']], r['Change category'], r['Description']) for r in taxonomy]
    common = sorted([r for r in rows if r['release_share'] >= 0.1], key=lambda r: (-r['release_pairs'], r['code']))
    rare = [r for r in rows if r['release_share'] < 0.1]
    grouped = list(common)
    if rare:
        grouped.append(count([r['code'] for r in rare], 'Other',
                             'Categories below 10% of changed release pairs.'))
    return rows, grouped


def summarize_breadth(pairs):
    rows = []
    for entity in ('file', 'procedure'):
        values = [pair[entity] for pair in pairs.values()]
        applicable = [v for v in values if v['share'] is not None]
        affected = sum(v['affected'] for v in values)
        total = sum(v['total'] for v in values)
        rows.append(dict(entity=entity, affected=affected, total=total,
            pooled_share=affected / total if total else None,
            min_affected=min((v['affected'] for v in applicable), default=None),
            median_affected=median(v['affected'] for v in values) if values else None,
            max_affected=max((v['affected'] for v in applicable), default=None),
            median_share=median(v['share'] for v in applicable) if applicable else None,
            pairs_affecting_multiple=sum(v['affected'] >= 2 for v in values),
            multiple_pair_share=share(sum(v['affected'] >= 2 for v in values), len(values)),
            zero_affected_pairs=sum(v['affected'] == 0 for v in values),
            no_applicable_locations=sum(v['total'] == 0 for v in values)))
    return rows


def calculate(releases, pairs, matches, screens, annotations, taxonomy, rows):
    # The CSV keeps resolved screening outcomes; omitted matches remain unresolved.
    screens = list(screens)
    recorded = {row['match_id'] for row in screens}
    screens.extend(dict(case_id='unresolved:' + key, match_id=key,
                        screen_status='unresolved', review_outcome='')
                   for key in sorted(set(matches) - recorded))
    pairs = {r['pair_id']: r for r in pairs}
    selected = check_selection(screens, annotations, matches, pairs, releases)
    changed = {cid for cid, r in annotations.items() if r['Outcome'] == 'change'}
    locations = location_sets(rows, changed)
    measures = aggregate_locations(locations, selected)
    categories, grouped = category_counts(annotations, selected, taxonomy)
    breadth = summarize_breadth(measures)
    release_rows = []
    for pair_id, values in sorted(measures.items()):
        pair = pairs[pair_id]
        old, new = releases[pair['old_release_id']], releases[pair['new_release_id']]
        for entity, m in values.items():
            release_rows.append(dict(pair_id=pair_id, repository=old['repository'], old_tag=old['tag'],
                new_tag=new['tag'], entity=entity, affected=m['affected'], total=m['total'], share=m['share']))
    summary = dict(
        population=dict(binding_pairs=len(selected), release_pairs=len({r['pair_id'] for r in selected.values()})),
        screening=dict(Counter(row['screen_status'] for row in screens)),
        review_outcomes=dict(Counter(row['review_outcome'] for row in screens if row['review_outcome'])),
        reviewed_binding_pairs=len(annotations), changing_binding_pairs=len(changed),
        changing_binding_share=share(len(changed), len(annotations)),
        changed_release_pairs=len(measures), changed_repositories=len({r['repository'] for r in release_rows}),
        path_surface=dict(Counter(combined_surface([m['surface'] for m in values.values()])
                                  for values in measures.values())),
        frequency_scope='Valid reviewed binding pairs with code edits and unchanged repository-wide PTM-ID counts.')
    tables = dict(integration_categories=categories, integration_categories_grouped=grouped,
                  integration_breadth=breadth, integration_release_pairs=release_rows)
    return tables, summary, locations, measures


def write_results(path, tables):
    if path.exists():
        raise FileExistsError('Results workbook already exists')
    book = Workbook()
    book.remove(book.active)
    for name in ('breadth', 'categories_grouped', 'categories', 'release_pairs'):
        rows = tables['integration_' + name]
        fields = (list(rows[0]) if rows else
                  ['pair_id', 'repository', 'old_tag', 'new_tag', 'entity', 'affected', 'total', 'share'])
        sheet = book.create_sheet(name)
        sheet.append(fields)
        for row in rows:
            sheet.append([row[field] for field in fields])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            width = max(len(str(cell.value or '')) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(70, max(12, width + 2))
            for cell in column:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    book.save(path)
    book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--database', type=Path, required=True)
    parser.add_argument('--inputs', type=Path, default=HERE / 'inputs')
    parser.add_argument('--output', type=Path, required=True, help='Output directory; existing integration results are not overwritten')
    args = parser.parse_args()
    if any((args.output / name).exists() for name in ('integration_results.xlsx', 'integration_summary.json')):
        raise FileExistsError('Output directory already contains integration results')
    files = {name: args.inputs / name for name in ('raw_ptm_counts.csv', 'integration_screening.csv',
                                                  'integration_change_annotations.xlsx', 'integration_locations.csv')}
    hashes = {name: sha256(path) for name, path in dict(files, database=args.database).items()}
    releases, pairs, _ = read_database(args.database)
    counts = read_counts(files['raw_ptm_counts.csv'], releases)
    matches = read_matches(args.database, pairs, counts)
    annotations, taxonomy = read_annotations(files['integration_change_annotations.xlsx'])
    tables, summary, _, _ = calculate(releases, pairs, matches,
        read_csv(files['integration_screening.csv']), annotations, taxonomy,
        read_csv(files['integration_locations.csv']))
    if hashes != {name: sha256(path) for name, path in dict(files, database=args.database).items()}:
        raise ValueError('Inputs changed during analysis')
    args.output.mkdir(parents=True, exist_ok=True)
    write_results(args.output / 'integration_results.xlsx', tables)
    with (args.output / 'integration_summary.json').open('x') as stream:
        json.dump(summary, stream, indent=2)
        stream.write('\n')
    print(f"Wrote {summary['changing_binding_pairs']} integration changes from {summary['reviewed_binding_pairs']} valid reviews.")


if __name__ == '__main__':
    main()
