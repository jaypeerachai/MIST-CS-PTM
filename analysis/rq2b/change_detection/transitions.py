"""Generate update and migration classifications from reviewed PTM-ID components."""

import argparse
from itertools import permutations
from pathlib import Path

from openpyxl import load_workbook

from common import write_csv

HERE = Path(__file__).resolve().parent
FIELDS = ('ptm_id', 'provider', 'product_lineage', 'functional_variant',
          'generation', 'version', 'revision', 'release_status', 'variant_suffix')


def read_components(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book['components']
        rows = sheet.iter_rows()
        header = [cell.value for cell in next(rows)]
        if len(header) != len(set(header)) or not set(FIELDS) <= set(header):
            raise ValueError('PTM component sheet has missing or repeated columns')
        components = []
        seen = set()
        for cells in rows:
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == 'f' for cell in cells):
                raise ValueError('PTM components must be reviewed values, not formulas')
            row = {name: '' if cell.value is None else str(cell.value).strip()
                   for name, cell in zip(header, cells)}
            if any(not row[field] for field in FIELDS[:4]) or row['ptm_id'] in seen:
                raise ValueError('Missing identity components or repeated PTM ID')
            seen.add(row['ptm_id'])
            components.append(row)
        if not components:
            raise ValueError('PTM component sheet is empty')
        return sorted(components, key=lambda row: row['ptm_id'])
    finally:
        book.close()


def classify(old, new):
    if old['ptm_id'] == new['ptm_id']:
        raise ValueError('A transition requires different PTM IDs')
    # Generation, version, revision, release status, and suffix changes stay within an update.
    for field in ('provider', 'product_lineage', 'functional_variant'):
        if old[field] != new[field]:
            return 'migration'
    return 'update'


def build_transitions(components):
    return {(old['ptm_id'], new['ptm_id']): classify(old, new)
            for old, new in permutations(components, 2)}


def write_transitions(path, transitions):
    rows = [dict(old_ptm_id=old, new_ptm_id=new, category=category)
            for (old, new), category in sorted(transitions.items())]
    write_csv(path, rows, ['old_ptm_id', 'new_ptm_id', 'category'])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--components', type=Path, default=HERE.parent / 'inputs/ptm_id_components.xlsx')
    parser.add_argument('--output', type=Path, required=True, help='New model_transition.csv file')
    args = parser.parse_args()
    components = read_components(args.components)
    transitions = build_transitions(components)
    write_transitions(args.output, transitions)
    print(f'Wrote {len(transitions):,} transitions from {len(components)} PTM IDs.')


if __name__ == '__main__':
    main()
