"""Input signatures and local-file handling around the unchanged TSE analyzer."""
from __future__ import annotations
import importlib.util
import re
import sys
import types
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from openpyxl import load_workbook

UPSTREAM_PATH = Path(__file__).resolve().parent / 'upstream/tse/static_analysis_fp_mapping.py'
TSE_FP_PATH_TOKENS = (
    "example",
    "examples",
    "lib/site-packages",
    "demo",
    "demos",
    "tutorial",
    "tutorials",
    "sample",
    "samples",
    ".venv",
    "environment",
    "environments",
    "env",
    "envs",
)


@dataclass(frozen=True)
class TSERule:
    import_signature: str
    call_signature: str
    model_loader: str
    chain_suffix: str
    terminal_call: str
    model_args: str
    round_status: str


@dataclass(frozen=True)
class TSEMatch:
    import_signature: str
    call_signature: str
    import_origin: str
    call_name: str
    call_line_number: int
    param_value: str
    param_line_number: int
    normalized_param_value: str
    model_loader: str


def load_upstream_module() -> types.ModuleType:
    install_utility_shims()
    spec = importlib.util.spec_from_file_location("tse_upstream_static_analysis_fp_mapping", UPSTREAM_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load upstream TSE script from {UPSTREAM_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules["tse_upstream_static_analysis_fp_mapping"] = module
    spec.loader.exec_module(module)
    return module


def install_utility_shims() -> None:
    """Satisfy upstream database/config imports without running its main block."""
    if "requests" not in sys.modules:
        requests_shim = types.ModuleType("requests")

        class _Session:
            def get(self, *_args: Any, **_kwargs: Any) -> Any:
                raise RuntimeError("requests shim does not support network access")

        def _get(*_args: Any, **_kwargs: Any) -> Any:
            raise RuntimeError("requests shim does not support network access")

        requests_shim.Session = _Session
        requests_shim.get = _get
        sys.modules["requests"] = requests_shim

    utilities = sys.modules.setdefault("utilities", types.ModuleType("utilities"))
    utilities.__path__ = []  # type: ignore[attr-defined]

    db_config = types.ModuleType("utilities.DBConfig")
    db_config.DatabaseConfig = type("DatabaseConfig", (), {})
    sys.modules["utilities.DBConfig"] = db_config

    db_schema = types.ModuleType("utilities.DBSchema")
    db_schema.DBTableNames = type("DBTableNames", (), {})
    db_schema.DBFieldNames = type("DBFieldNames", (), {})
    sys.modules["utilities.DBSchema"] = db_schema

    hf_config = types.ModuleType("utilities.HFConfig")
    hf_config.HuggingFaceConfig = type("HuggingFaceConfig", (), {})
    sys.modules["utilities.HFConfig"] = hf_config

    raw_file_config = types.ModuleType("utilities.RawFileConfig")
    raw_file_config.RawFileConfig = type("RawFileConfig", (), {})
    sys.modules["utilities.RawFileConfig"] = raw_file_config


def load_rules(path: Path) -> list[TSERule]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["model_loading"]
    rows = ws.iter_rows(values_only=True)
    headers = [normalize_header(value) for value in next(rows)]
    pos = {header: index for index, header in enumerate(headers) if header}
    required = ["import origin", "model loader", "terminal_call", "chain_suffix", "model argument(s)", "round status"]
    missing = [header for header in required if header not in pos]
    if missing:
        raise ValueError(f"Missing reuse codebook columns: {missing}")

    rules: list[TSERule] = []
    seen = set()
    for values in rows:
        row = {header: clean_cell(values[index] if index < len(values) else "") for header, index in pos.items()}
        import_signature = row.get("import origin", "")
        terminal_call = row.get("terminal_call", "")
        if not import_signature or not terminal_call:
            continue
        rule = TSERule(
            import_signature=import_signature,
            call_signature=terminal_call,
            model_loader=row.get("model loader", ""),
            chain_suffix=row.get("chain_suffix", ""),
            terminal_call=terminal_call,
            model_args=row.get("model argument(s)", ""),
            round_status=row.get("round status", ""),
        )
        key = (rule.import_signature, rule.call_signature, rule.model_loader, rule.chain_suffix)
        if key in seen:
            continue
        seen.add(key)
        rules.append(rule)
    return rules


def group_rules_by_origin(rules: Iterable[TSERule]) -> dict[str, list[TSERule]]:
    grouped: dict[str, list[TSERule]] = defaultdict(list)
    for rule in rules:
        grouped[rule.import_signature].append(rule)
    return dict(grouped)


def run_tse_on_file(
    *,
    analyzer_cls: type,
    repo_path: Path,
    rel_path: str,
    rules_by_origin: dict[str, list[TSERule]],
    unwrap_quotes: Any,
    looks_like_local_path: Any,
    non_result_cache: set[str],
) -> tuple[list[TSEMatch], str]:
    if path_filtered(rel_path):
        return [], "tse_path_filter"
    source_path = repo_path / rel_path
    if not source_path.exists():
        return [], f"missing_file:{source_path}"

    analyzer = analyzer_cls(str(source_path))
    try:
        analyzer.load_and_parse()
    except (IndentationError, TabError, SyntaxError) as exc:
        return [], f"parse_error:{type(exc).__name__}:{exc}"
    except Exception as exc:
        return [], f"read_or_parse_error:{type(exc).__name__}:{exc}"

    matches: list[TSEMatch] = []
    for import_signature, rules in rules_by_origin.items():
        for rule in rules:
            try:
                raw_matches = analyzer.analyze(
                    import_signature=import_signature,
                    call_signature=rule.call_signature,
                )
            except Exception as exc:
                return matches, f"tse_analyze_error:{type(exc).__name__}:{exc}"
            matches.extend(
                convert_matches(
                    rule=rule,
                    raw_matches=raw_matches,
                    unwrap_quotes=unwrap_quotes,
                    looks_like_local_path=looks_like_local_path,
                    non_result_cache=non_result_cache,
                )
            )
    return dedupe_matches(matches), "ok"


def convert_matches(
    *,
    rule: TSERule,
    raw_matches: list[dict[str, Any]],
    unwrap_quotes: Any,
    looks_like_local_path: Any,
    non_result_cache: set[str],
) -> list[TSEMatch]:
    output: list[TSEMatch] = []
    for raw in raw_matches:
        if raw.get("import_origin") != rule.import_signature:
            continue
        call_name = str(raw.get("name", ""))
        call_line = parse_int(raw.get("lineno"))
        for param in raw.get("resolved_params", []):
            raw_value = param.get("value")
            value = "" if raw_value is None else str(raw_value)
            value = unwrap_quotes(value).strip()
            if "models/" in value:
                value = value.split("/", 1)[-1]
            normalized = value.lower()
            if not normalized or normalized in non_result_cache:
                continue
            if looks_like_local_path(value):
                continue
            output.append(
                TSEMatch(
                    import_signature=rule.import_signature,
                    call_signature=rule.call_signature,
                    import_origin=str(raw.get("import_origin", "")),
                    call_name=call_name,
                    call_line_number=call_line,
                    param_value=value,
                    param_line_number=parse_int(param.get("assign_lineno", call_line)),
                    normalized_param_value=normalized,
                    model_loader=rule.model_loader,
                )
            )
    return output


def dedupe_matches(matches: list[TSEMatch]) -> list[TSEMatch]:
    seen = set()
    output = []
    for match in sorted(
        matches,
        key=lambda item: (
            item.call_line_number,
            item.param_line_number,
            item.call_name,
            item.normalized_param_value,
        ),
    ):
        key = (
            match.import_signature,
            match.call_signature,
            match.call_line_number,
            match.param_line_number,
            match.call_name,
            match.normalized_param_value,
        )
        if key in seen:
            continue
        seen.add(key)
        output.append(match)
    return output


def model_matches_param(model_id: str, normalized_param_value: str) -> bool:
    full = model_id.strip().strip("'\"").lower()
    bare = full.rsplit("/", 1)[-1]
    value = normalized_param_value.strip().strip("'\"").lower()
    return value == full or value == bare


def path_filtered(rel_path: str) -> bool:
    lower = rel_path.lower()
    return any(token in lower for token in TSE_FP_PATH_TOKENS)


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def clean_cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_int(value: Any) -> int:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return 0
