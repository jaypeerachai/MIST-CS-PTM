"""Input signatures and argument matching around the unchanged PeaTMOSS analyzer."""
from __future__ import annotations
import importlib.util
import json
import re
import sys
import traceback
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from openpyxl import load_workbook

UPSTREAM_DIR = Path(__file__).resolve().parent / 'upstream/peatmoss'


@dataclass(frozen=True)
class PeaTMOSSRule:
    import_origin: str
    class_to_check: str
    method_to_check: str
    model_loader: str
    chain_suffix: str
    terminal_call: str
    model_args: str
    round_status: str


@dataclass(frozen=True)
class LoaderMatch:
    import_origin: str
    model_loader: str
    chain_suffix: str
    terminal_call: str
    class_to_check: str
    method_to_check: str
    call_name: str
    line_number: int
    params: str
    model_args: str


def install_upstream_mnode() -> None:
    """Make PeaTMOSS' customized MNode satisfy ExtractSource's import."""
    mnode_path = UPSTREAM_DIR / "scalpel.core.mnode.py"
    spec = importlib.util.spec_from_file_location("scalpel.core.mnode", mnode_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load upstream MNode from {mnode_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules["scalpel.core.mnode"] = module
    spec.loader.exec_module(module)


def load_extract_source_class() -> type:
    extract_path = UPSTREAM_DIR / "ExtractSource.py"
    spec = importlib.util.spec_from_file_location("peatmoss_upstream_extract_source", extract_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load upstream ExtractSource from {extract_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules["peatmoss_upstream_extract_source"] = module
    spec.loader.exec_module(module)
    return module.ExtractSource


def load_peatmoss_rules(path: Path) -> list[PeaTMOSSRule]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["model_loading"]
    rows = ws.iter_rows(values_only=True)
    headers = [normalize_header(value) for value in next(rows)]
    pos = {header: index for index, header in enumerate(headers) if header}
    required = ["import origin", "model loader", "terminal_call", "chain_suffix", "model argument(s)", "round status"]
    missing = [header for header in required if header not in pos]
    if missing:
        raise ValueError(f"Missing reuse codebook columns: {missing}")

    rules: list[PeaTMOSSRule] = []
    seen = set()
    for values in rows:
        row = {header: clean_cell(values[index] if index < len(values) else "") for header, index in pos.items()}
        import_origin = row.get("import origin", "")
        terminal_call = row.get("terminal_call", "")
        model_loader = row.get("model loader", "")
        chain_suffix = row.get("chain_suffix", "")
        if not import_origin or not terminal_call:
            continue
        class_to_check, method_to_check = peatmoss_signature_parts(
            model_loader=model_loader,
            chain_suffix=chain_suffix,
            terminal_call=terminal_call,
        )
        key = (import_origin, class_to_check, method_to_check, model_loader, chain_suffix)
        if key in seen:
            continue
        seen.add(key)
        rules.append(
            PeaTMOSSRule(
                import_origin=import_origin,
                class_to_check=class_to_check,
                method_to_check=method_to_check,
                model_loader=model_loader,
                chain_suffix=chain_suffix,
                terminal_call=terminal_call,
                model_args=row.get("model argument(s)", ""),
                round_status=row.get("round status", ""),
            )
        )
    return rules


def peatmoss_signature_parts(*, model_loader: str, chain_suffix: str, terminal_call: str) -> tuple[str, str]:
    method = terminal_call.strip() or terminal_from_text(chain_suffix) or terminal_from_text(model_loader)
    chain = chain_suffix.strip() or model_loader.strip()
    parts = [part for part in chain.split(".") if part]
    class_to_check = ""
    if method and method in parts:
        index = len(parts) - 1 - list(reversed(parts)).index(method)
        if index > 0:
            class_to_check = parts[index - 1]
        else:
            class_to_check = method
    elif len(parts) >= 2:
        class_to_check = parts[-2]
    elif len(parts) == 1:
        class_to_check = parts[0]
    return class_to_check, method


def group_rules_by_origin(rules: Iterable[PeaTMOSSRule]) -> dict[str, list[PeaTMOSSRule]]:
    grouped: dict[str, list[PeaTMOSSRule]] = defaultdict(list)
    for rule in rules:
        grouped[rule.import_origin].append(rule)
    return dict(grouped)


def run_peatmoss_on_file(
    *,
    extract_source_cls: type,
    repo_path: Path,
    rel_path: str,
    rules_by_origin: dict[str, list[PeaTMOSSRule]],
) -> tuple[list[LoaderMatch], str]:
    source_path = repo_path / rel_path
    if not source_path.exists():
        return [], f"missing_file:{source_path}"
    try:
        source = source_path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        return [], f"read_error:{type(exc).__name__}:{exc}"

    matches: list[LoaderMatch] = []
    for import_origin, rules in rules_by_origin.items():
        extractor = extract_source_cls(source)
        extractor.set_blob(f"{repo_path.name}:{rel_path}:{import_origin}")
        extractor.set_importname(import_origin)
        extractor.set_targets_info(
            [
                {
                    "class_to_check": rule.class_to_check,
                    "method_to_check": rule.method_to_check,
                }
                for rule in rules
            ]
        )
        try:
            raw_matches = extractor.fetch_match_called()
        except Exception:
            return matches, "peatmoss_error:" + traceback.format_exc(limit=1).strip()
        matches.extend(convert_matches(import_origin, rules, raw_matches))
    return dedupe_matches(matches), "ok"


def convert_matches(import_origin: str, rules: list[PeaTMOSSRule], raw_matches: list[dict[str, Any]]) -> list[LoaderMatch]:
    output: list[LoaderMatch] = []
    for raw in raw_matches:
        call_name = str(raw.get("name", ""))
        params = json.dumps(raw.get("params", ""), ensure_ascii=False, sort_keys=True, default=str)
        line_number = parse_int(raw.get("lineno"))
        for rule in matching_rules_for_call(rules, call_name):
            output.append(
                LoaderMatch(
                    import_origin=import_origin,
                    model_loader=rule.model_loader,
                    chain_suffix=rule.chain_suffix,
                    terminal_call=rule.terminal_call,
                    class_to_check=rule.class_to_check,
                    method_to_check=rule.method_to_check,
                    call_name=call_name,
                    line_number=line_number,
                    params=params,
                    model_args=rule.model_args,
                )
            )
    return output


def matching_rules_for_call(rules: list[PeaTMOSSRule], call_name: str) -> list[PeaTMOSSRule]:
    parts = call_name.split(".")
    output = []
    for rule in rules:
        if rule.method_to_check not in parts[1:]:
            continue
        index = parts.index(rule.method_to_check)
        if rule.class_to_check == rule.method_to_check:
            output.append(rule)
        elif not rule.class_to_check and index:
            output.append(rule)
        elif index > 0 and parts[index - 1] == rule.class_to_check:
            output.append(rule)
    return output or [rule for rule in rules if rule.method_to_check in parts[1:]]


def dedupe_matches(matches: list[LoaderMatch]) -> list[LoaderMatch]:
    seen = set()
    output = []
    for match in sorted(matches, key=lambda item: (item.line_number, item.call_name, item.model_loader)):
        # PeaTMOSS returns call-site matches, not codebook-rule matches. Our
        # richer codebook can map several signatures to the same call; collapse
        # those duplicates so evidence remains call-site granular.
        key = (match.import_origin, match.call_name, match.line_number, match.params)
        if key in seen:
            continue
        seen.add(key)
        output.append(match)
    return output


def extracted_model_arg_values(match: LoaderMatch) -> list[str]:
    values = extracted_param_values(match.params)
    if not values:
        return []
    # PeaTMOSS' Scalpel call visitor records positional and keyword values in
    # call order but does not preserve keyword names. The codebook field is
    # still retained in outputs and used as a guard that the signature has a
    # model-bearing argument in the original schema.
    if not match.model_args:
        return []
    return [normalize_model_arg_value(value) for value in values if normalize_model_arg_value(value)]


def extracted_param_values(params: str) -> list[str]:
    try:
        parsed = json.loads(params)
    except json.JSONDecodeError:
        parsed = params
    if not isinstance(parsed, list):
        parsed = [parsed]
    values: list[str] = []
    for item in parsed:
        if item is None or isinstance(item, (bool, int, float)):
            continue
        value = str(item).strip()
        if value and value.lower() not in {"unknown", "dict", "list", "tuple", "set", "expr"}:
            values.append(value)
    return values


def normalize_model_arg_value(value: str) -> str:
    output = value.strip().strip("'\"").strip()
    if not output:
        return ""
    output = re.sub(r"^https?://huggingface\.co/", "", output, flags=re.IGNORECASE)
    output = output.split("?", 1)[0].split("#", 1)[0].strip("/")
    if output.lower().startswith("models/"):
        output = output.split("/", 1)[1]
    return output.lower()


def model_matches_arg(model_id: str, normalized_arg_value: str) -> bool:
    full = normalize_model_arg_value(model_id)
    bare = full.rsplit("/", 1)[-1]
    value = normalized_arg_value.strip().strip("'\"").lower()
    return bool(value) and (value == full or value == bare)


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def clean_cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def terminal_from_text(value: str) -> str:
    return value.rsplit(".", 1)[-1].strip()


def parse_int(value: Any) -> int:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return 0
