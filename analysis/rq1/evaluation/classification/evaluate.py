"""Score PTM reuse classification against the reviewed cases."""

import argparse
import csv
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import random
from urllib.parse import unquote, urlsplit

from openpyxl import load_workbook
import numpy as np
from sklearn.metrics import accuracy_score, confusion_matrix, precision_recall_fscore_support

ROOT = Path(__file__).resolve().parent
DATASETS = {'benchmark': ROOT / 'benchmark.xlsx',
            'unseen_holdout': ROOT / 'unseen_holdout.xlsx'}
LABELS = {'real_reuse': True, 'non_reuse': False}
METHODS = {'mist': 'MIST', 'peatmoss': 'Jiang et al.', 'tse': 'Banyongrakkul et al.'}


def rounded(value):
    return Decimal(str(value)).quantize(Decimal('0.001'), rounding=ROUND_HALF_UP)


def read_csv(path):
    with Path(path).open(newline='', encoding='utf-8-sig') as handle:
        return list(csv.DictReader(handle))


def read_sheet(book, name, key):
    rows = iter(book[name].values)
    headers = next(rows)
    records = {}
    for values in rows:
        if not any(value is not None for value in values):
            continue
        row = dict(zip(headers, values))
        identity = row[key]
        if not identity or identity in records:
            raise ValueError(f'Missing or duplicate {key} in {name}: {identity}')
        records[identity] = row
    return records


def load_dataset(name):
    book = load_workbook(DATASETS[name], read_only=True, data_only=True)
    try:
        annotations = read_sheet(book, 'annotations', 'case_id')
        predictions = read_sheet(book, 'predictions', 'case_id')
        repositories = read_sheet(book, 'repositories', 'repository')
    finally:
        book.close()
    if set(annotations) != set(predictions):
        raise ValueError('Annotations and predictions have different case IDs')
    for case_id, row in annotations.items():
        parts = unquote(urlsplit(row['source_url']).path).lstrip('/').split('/')
        if len(parts) < 5 or parts[2] != 'blob' or '/'.join(parts[:2]) != row['repository']:
            raise ValueError(f'Invalid fixed source URL: {case_id}')
        row['commit'], row['file_path'] = parts[3], '/'.join(parts[4:])
        if row['commit'] != repositories[row['repository']]['commit']:
            raise ValueError(f'Commit mismatch: {case_id}')
        if row['label'] not in LABELS:
            raise ValueError(f'Unknown label: {case_id}')
        if any(predictions[case_id].get(method) not in LABELS for method in METHODS):
            raise ValueError(f'Missing prediction: {case_id}')
    return annotations, predictions


def load_repositories(dataset='all', repository=None):
    jobs = []
    for name, path in DATASETS.items():
        if dataset not in ('all', name):
            continue
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = read_sheet(book, 'repositories', 'repository')
        finally:
            book.close()
        for row in rows.values():
            if repository is None or row['repository'] == repository:
                jobs.append(dict(dataset=name, repository=row['repository'], commit=row['commit']))
    return sorted(jobs, key=lambda row: (row['dataset'], row['repository'], row['commit']))


def load_predictions(path):
    result = {}
    for row in read_csv(path):
        if row['case_id'] in result or row['prediction'] not in LABELS:
            raise ValueError(f'Duplicate case or unknown prediction: {row}')
        result[row['case_id']] = row['prediction']
    return result


def score(pairs):
    gold, predicted = zip(*pairs)
    tn, fp, fn, tp = confusion_matrix(gold, predicted, labels=[False, True]).ravel()
    precision, recall, f1, _ = precision_recall_fscore_support(
        gold, predicted, average='binary', pos_label=True, zero_division=0)
    return dict(TP=int(tp), FP=int(fp), FN=int(fn), TN=int(tn),
                precision=float(precision), recall=float(recall), f1=float(f1),
                accuracy=float(accuracy_score(gold, predicted)))


def classification(annotations, predictions):
    repositories = defaultdict(list)
    for case_id, row in annotations.items():
        repositories[row['repository']].append(case_id)
    results = []
    for level in ('occurrence', 'repository'):
        groups = [[case_id] for case_id in annotations] if level == 'occurrence' else repositories.values()
        for method in METHODS:
            pairs = [(any(LABELS[annotations[k]['label']] for k in group),
                      any(LABELS[predictions[k][method]] for k in group)) for group in groups]
            results.append(dict(level=level, method=method, **score(pairs)))
    return results


def bootstrap_comparison(annotations, predictions, resamples=10_000, seed=20260803):
    if resamples <= 0 or not annotations:
        raise ValueError('Bootstrap needs cases and a positive number of resamples')
    if set(annotations) != set(predictions):
        raise ValueError('Annotations and predictions have different case IDs')
    by_repository = defaultdict(list)
    for case_id, row in annotations.items():
        by_repository[row['repository']].append(case_id)
    repositories = sorted(by_repository)
    case_ids = list(annotations)
    pairs = {}
    for level in ('occurrence', 'repository'):
        groups = [[case_id] for case_id in case_ids] if level=='occurrence' else [by_repository[r] for r in repositories]
        pairs[level] = {}
        for method in METHODS:
            gold = [any(LABELS[annotations[k]['label']] for k in group) for group in groups]
            predicted = [any(LABELS[predictions[k][method]] for k in group) for group in groups]
            pairs[level][method] = (gold, predicted)
    gains = {(level, baseline, metric): [] for level in pairs
             for baseline in METHODS if baseline!='mist' for metric in ('recall','f1')}
    rng = random.Random(seed)
    for _ in range(resamples):
        counts = Counter(rng.choice(repositories) for _ in repositories)
        # A weight repeats the whole repository, including all its occurrences.
        weights = {'occurrence': [counts[annotations[k]['repository']] for k in case_ids],
                   'repository': [counts[r] for r in repositories]}
        for level, methods in pairs.items():
            scores = {}
            for method, (gold, predicted) in methods.items():
                _, recall, f1, _ = precision_recall_fscore_support(
                    gold, predicted, average='binary', pos_label=True,
                    sample_weight=weights[level], zero_division=0)
                scores[method] = {'recall': recall, 'f1': f1}
            for baseline in METHODS:
                if baseline!='mist':
                    for metric in ('recall','f1'):
                        gains[level,baseline,metric].append(100*(scores['mist'][metric]-scores[baseline][metric]))
    points = {(r['level'],r['method']):r for r in classification(annotations,predictions)}
    results = []
    for (level,baseline,metric), values in gains.items():
        low, high = np.percentile(values, [2.5,97.5], method='linear')
        results.append(dict(level=level,baseline=baseline,metric=metric,
            gain_pp=100*(points[level,'mist'][metric]-points[level,baseline][metric]),
            ci_low_pp=float(low),ci_high_pp=float(high),resamples=resamples,seed=seed))
    return results


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dataset', choices=[*DATASETS, 'all'], default='all')
    parser.add_argument('--predictions', type=Path, help='score a new run: CSV with case_id,prediction')
    parser.add_argument('--prediction-method', choices=list(METHODS), default='mist',
                        help='method supplied by --predictions (default: mist)')
    parser.add_argument('--output', type=Path, help='write scores to a new CSV')
    parser.add_argument('--bootstrap', action='store_true', help='compare MIST with both baselines on the holdout')
    parser.add_argument('--resamples', type=int, default=10_000, help='bootstrap draws (default: 10000)')
    parser.add_argument('--seed', type=int, default=20260803, help='bootstrap random seed')
    parser.add_argument('--bootstrap-output', type=Path, help='write bootstrap comparisons to a new CSV')
    args = parser.parse_args()
    if args.prediction_method != 'mist' and not args.predictions:
        parser.error('--prediction-method requires --predictions')
    if args.bootstrap and (args.dataset=='benchmark' or args.resamples<=0):
        parser.error('Bootstrap comparisons require the holdout and a positive number of resamples')
    if args.bootstrap_output and not args.bootstrap:
        parser.error('--bootstrap-output requires --bootstrap')
    if args.output and args.bootstrap_output and args.output.resolve()==args.bootstrap_output.resolve():
        parser.error('Use separate output paths for classification and bootstrap results')
    names = list(DATASETS) if args.dataset == 'all' else [args.dataset]
    replacement = load_predictions(args.predictions) if args.predictions else None
    if replacement is not None:
        known = {k for name in DATASETS for k in load_dataset(name)[0]}
        if set(replacement) - known:
            raise ValueError('Predictions contain unknown case IDs')
    results, intervals = [], []
    print(f'{METHODS[args.prediction_method]} predictions:', args.predictions or 'workbook predictions sheets')
    for name in names:
        annotations, predictions = load_dataset(name)
        if replacement is not None:
            missing = set(annotations) - set(replacement)
            if missing:
                raise ValueError(f'Missing {args.prediction_method} predictions for {len(missing)} cases; not counted as negatives')
            for case_id in annotations:
                predictions[case_id][args.prediction_method] = replacement[case_id]
        print(f'\n{name}: {len(annotations)} occurrences, {len({r["repository"] for r in annotations.values()})} repositories')
        print(f'{"Level":12} {"Method":22} {"TP":>4} {"FP":>4} {"FN":>4} {"TN":>4} {"P":>6} {"R":>6} {"F1":>6}')
        for row in classification(annotations, predictions):
            results.append(dict(dataset=name, **row))
            print(f'{row["level"]:12} {METHODS[row["method"]]:22} ' +
                  ' '.join(f'{row[k]:4}' for k in ('TP','FP','FN','TN')) + ' ' +
                  ' '.join(f'{rounded(row[k]):6}' for k in ('precision','recall','f1')))
        if name == 'benchmark':
            for scope in ('local','non_local'):
                selected = [k for k,r in annotations.items() if r.get('binding_locality') == scope]
                counts = ', '.join(f'{METHODS[m]} {sum(LABELS[predictions[k][m]] for k in selected)}/{len(selected)}' for m in METHODS)
                print(f'{scope} binding recovery: {counts}')
        if args.bootstrap and name=='unseen_holdout':
            print(f'\nPaired repository bootstrap: {args.resamples:,} draws, seed {args.seed}',flush=True)
            intervals = bootstrap_comparison(annotations,predictions,args.resamples,args.seed)
            print('MIST gains in percentage points, with 95% percentile intervals')
            print(f'{"Level":12} {"Baseline":22} {"Metric":8} {"Gain":>7} {"95% interval":>17}')
            for row in intervals:
                print(f'{row["level"]:12} {METHODS[row["baseline"]]:22} {row["metric"]:8} '
                      f'{row["gain_pp"]:7.1f} [{row["ci_low_pp"]:6.1f}, {row["ci_high_pp"]:6.1f}]')
    if args.output:
        with args.output.open('x', newline='', encoding='utf-8') as handle:
            writer = csv.DictWriter(handle, fieldnames=list(results[0]))
            writer.writeheader()
            writer.writerows(results)
    if args.bootstrap_output:
        with args.bootstrap_output.open('x',newline='',encoding='utf-8') as handle:
            writer = csv.DictWriter(handle,fieldnames=list(intervals[0]))
            writer.writeheader()
            writer.writerows(intervals)


if __name__ == '__main__':
    main()
